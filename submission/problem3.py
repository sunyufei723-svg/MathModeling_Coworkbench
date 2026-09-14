"""问题三：固定半径、附录3变物性热湿耦合，守恒径向有限体积 + 联合状态自适应 BDF，求全截面达标停机的连续临界时间。

温度场与水分场通过附录3物性（ρ、cp、k、D 均随含水率 C 变化）双向强耦合，合并成联合状态 [T_0..T_N,C_0..C_N]
整体推进；界面 k、D 用调和平均；以全截面最大含水率事件函数 g(t)=max(C)-0.15 定位连续达标时刻 t*。
main 内自动完成论文「模型检验与稳定性分析」章引用的全部验证：空间网格逐级加密（观察收敛阶 p、Richardson
外推、调和平均-网格序列）、BDF 容差三档、后向欧拉-Picard 交叉对照、水分守恒残差与径向单调性、4h 后边界
敏感性（3×3 工程情景 / 传质系数 ±10% / 末值-30min-1h 均值延拓）。

直接运行（无需任何参数）：python problem3.py
  · 读同根目录 附件/附件1.xlsx（环境）与 附件/附件3/result3.xlsx（结果格式模板），读不到给提示；
  · 按模板格式把达标停机水分表写到同根目录 result3.xlsx，验证摘要写 verification3.json；
  · 终端分节打印主答案 t* + 网格收敛 + BDF 容差 + 双求解器交叉 + 守恒核验 + 边界敏感性。
  · 注意：本问验证含最细网格（Δr=0.00078125 cm，2561 节点）多次求解，完整运行约需数十分钟。
"""
from __future__ import annotations

import json
import time
from dataclasses import asdict, dataclass, replace
from math import ceil
from pathlib import Path
from typing import Callable

import numpy as np
import openpyxl
import pandas as pd
from scipy.integrate import solve_ivp
from scipy.linalg.lapack import dgtsv
from scipy.sparse import bmat, csc_matrix, diags


# ============================== 求解核心 · 物性与环境边界（附录3）==============================
@dataclass(frozen=True)
class EnvironmentSeries:
    """附件1给出的炉内边界序列；附件区间内只做分段线性插值。"""

    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def __post_init__(self) -> None:
        time = np.asarray(self.time_s, dtype=float)
        temperature = np.asarray(self.temperature_c, dtype=float)
        moisture = np.asarray(self.moisture, dtype=float)
        if not (time.ndim == temperature.ndim == moisture.ndim == 1):
            raise ValueError("environment arrays must be one-dimensional")
        if not (len(time) == len(temperature) == len(moisture)) or len(time) < 2:
            raise ValueError("environment arrays must have equal length >= 2")
        if np.any(np.diff(time) <= 0.0):
            raise ValueError("environment time must be strictly increasing")
        if not (np.isfinite(time).all() and np.isfinite(temperature).all() and np.isfinite(moisture).all()):
            raise ValueError("environment arrays must be finite")
        object.__setattr__(self, "time_s", time)
        object.__setattr__(self, "temperature_c", temperature)
        object.__setattr__(self, "moisture", moisture)

    @property
    def terminal_time_s(self) -> float:
        return float(self.time_s[-1])

    @property
    def terminal_temperature_c(self) -> float:
        return float(self.temperature_c[-1])

    @property
    def terminal_moisture(self) -> float:
        return float(self.moisture[-1])

    def interpolated_temperature(self, t_s: float) -> float:
        if t_s < self.time_s[0] or t_s > self.time_s[-1]:
            raise ValueError("interpolation time is outside the attachment interval")
        return float(np.interp(t_s, self.time_s, self.temperature_c))

    def interpolated_moisture(self, t_s: float) -> float:
        if t_s < self.time_s[0] or t_s > self.time_s[-1]:
            raise ValueError("interpolation time is outside the attachment interval")
        return float(np.interp(t_s, self.time_s, self.moisture))


@dataclass(frozen=True)
class PostBoundary:
    temperature_c: float
    moisture: float


def _safe_moisture(c: np.ndarray | float) -> np.ndarray:
    """只供物性计算使用；不修改求解器中的状态。"""

    return np.maximum(np.asarray(c, dtype=float), 1e-9)


def density(c: np.ndarray | float) -> np.ndarray:
    safe_c = _safe_moisture(c)
    return 650.0 + 128.0 * safe_c


def heat_capacity(c: np.ndarray | float) -> np.ndarray:
    safe_c = _safe_moisture(c)
    return 1450.0 + 2736.0 * safe_c / (safe_c + 1.0)


def thermal_conductivity(c: np.ndarray | float) -> np.ndarray:
    safe_c = _safe_moisture(c)
    return 0.21 + 0.38 * safe_c / (safe_c + 1.0)


def moisture_diffusivity(c: np.ndarray | float, temperature_c: np.ndarray | float) -> np.ndarray:
    safe_c = _safe_moisture(c)
    temperature_k = np.asarray(temperature_c, dtype=float) + 273.15
    if np.any(temperature_k <= 0.0):
        raise ValueError("temperature must be above absolute zero")
    return 2.4e-3 * np.exp(-0.45 / safe_c) * np.exp(-3850.0 / temperature_k)


def post_boundary_from_tail(environment: EnvironmentSeries, averaging_window_s: float | None = None) -> PostBoundary:
    if averaging_window_s is None:
        return PostBoundary(environment.terminal_temperature_c, environment.terminal_moisture)
    if averaging_window_s <= 0.0:
        raise ValueError("averaging_window_s must be positive")
    start = max(float(environment.time_s[0]), environment.terminal_time_s - averaging_window_s)
    interior = environment.time_s[(environment.time_s > start) & (environment.time_s <= environment.terminal_time_s)]
    times = np.concatenate([[start], interior])
    temperature = np.interp(times, environment.time_s, environment.temperature_c)
    moisture = np.interp(times, environment.time_s, environment.moisture)
    duration = float(times[-1] - times[0])
    if duration <= 0.0:
        raise ValueError("averaging window has zero duration")
    temperature_integral = float(np.sum(0.5 * (temperature[:-1] + temperature[1:]) * np.diff(times)))
    moisture_integral = float(np.sum(0.5 * (moisture[:-1] + moisture[1:]) * np.diff(times)))
    return PostBoundary(
        temperature_integral / duration,
        moisture_integral / duration,
    )


# ============================== 求解核心 · 径向 FVM 几何与界面离散 ==============================
@dataclass(frozen=True)
class RadialGeometry:
    radius_m: np.ndarray
    volumes_m3_per_m: np.ndarray
    interface_areas_m2_per_m: np.ndarray
    surface_area_m2_per_m: float
    dr_m: float

    @property
    def node_count(self) -> int:
        return int(self.radius_m.size)


def build_radial_geometry(radius_cm: float, dr_cm: float) -> RadialGeometry:
    if radius_cm <= 0.0 or dr_cm <= 0.0:
        raise ValueError("radius_cm and dr_cm must be positive")
    intervals = int(round(radius_cm / dr_cm))
    if intervals < 1 or not np.isclose(intervals * dr_cm, radius_cm, rtol=0.0, atol=1e-12):
        raise ValueError("radius_cm must be an integer multiple of dr_cm")
    radius_m = np.linspace(0.0, radius_cm / 100.0, intervals + 1)
    dr_m = dr_cm / 100.0
    left_faces = np.maximum(0.0, radius_m - dr_m / 2.0)
    right_faces = np.minimum(radius_m[-1], radius_m + dr_m / 2.0)
    volumes = np.pi * (right_faces**2 - left_faces**2)
    interface_areas = 2.0 * np.pi * (radius_m[:-1] + dr_m / 2.0)
    return RadialGeometry(
        radius_m=radius_m,
        volumes_m3_per_m=volumes,
        interface_areas_m2_per_m=interface_areas,
        surface_area_m2_per_m=float(2.0 * np.pi * radius_m[-1]),
        dr_m=dr_m,
    )


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    left = np.asarray(left, dtype=float)
    right = np.asarray(right, dtype=float)
    denominator = left + right
    result = np.zeros(np.broadcast_shapes(left.shape, right.shape), dtype=float)
    np.divide(2.0 * left * right, denominator, out=result, where=denominator != 0.0)
    return result


def interface_values(node_values: np.ndarray, mean: str) -> np.ndarray:
    node_values = np.asarray(node_values, dtype=float)
    if mean == "harmonic":
        return harmonic_mean(node_values[:-1], node_values[1:])
    if mean == "arithmetic":
        return 0.5 * (node_values[:-1] + node_values[1:])
    raise ValueError("interface_mean must be 'harmonic' or 'arithmetic'")


def internal_flux_numerator(
    values: np.ndarray,
    interface_transport: np.ndarray,
    geometry: RadialGeometry,
) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    numerator = np.zeros_like(values)
    flux = (
        np.asarray(interface_transport, dtype=float)
        * geometry.interface_areas_m2_per_m
        * (values[1:] - values[:-1])
        / geometry.dr_m
    )
    numerator[:-1] += flux
    numerator[1:] -= flux
    return numerator


def build_coupled_jacobian_sparsity(node_count: int) -> csc_matrix:
    if node_count < 2:
        raise ValueError("node_count must be at least two")
    tri = diags(
        [np.ones(node_count - 1), np.ones(node_count), np.ones(node_count - 1)],
        offsets=[-1, 0, 1],
        shape=(node_count, node_count),
        format="csc",
    )
    return bmat([[tri, tri], [tri, tri]], format="csc")


def sample_profiles(
    internal_radius_cm: np.ndarray,
    output_radius_cm: np.ndarray,
    profiles: np.ndarray,
) -> np.ndarray:
    internal_radius_cm = np.asarray(internal_radius_cm, dtype=float)
    output_radius_cm = np.asarray(output_radius_cm, dtype=float)
    profiles = np.asarray(profiles, dtype=float)
    indices = np.rint(output_radius_cm / (internal_radius_cm[1] - internal_radius_cm[0])).astype(int)
    if np.allclose(internal_radius_cm[indices], output_radius_cm, rtol=0.0, atol=1e-12):
        return profiles[:, indices]
    return np.vstack([np.interp(output_radius_cm, internal_radius_cm, row) for row in profiles])



# ============================== 求解核心 · 联合状态自适应 BDF 求解器（主求解器）==============================
@dataclass(frozen=True)
class SolverConfig:
    radius_cm: float = 2.0
    internal_dr_cm: float = 0.025
    output_dr_cm: float = 0.1
    output_interval_s: int = 60
    initial_temperature_c: float = 28.0
    initial_moisture: float = 2.55
    completion_threshold: float = 0.15
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8e-7
    interface_mean: str = "harmonic"
    relative_tolerance: float = 1e-8
    temperature_absolute_tolerance: float = 1e-8
    moisture_absolute_tolerance: float = 1e-10
    max_time_step_s: float = 60.0
    initial_time_limit_s: float = 72.0 * 3600.0
    fallback_time_limit_s: float = 96.0 * 3600.0
    extended_time_limit_s: float | None = None
    material_negative_tolerance: float = -1e-8
    monotonicity_tolerance: float = 1e-8


@dataclass(frozen=True)
class SolverDiagnostics:
    internal_node_count: int
    state_size: int
    function_evaluations: int
    jacobian_evaluations: int
    lu_decompositions: int
    phase1_end_time_s: float
    first_limit_reached_without_event: bool
    fallback_limit_reached_without_event: bool
    event_count: int
    event_direction: int
    minimum_temperature_c: float
    maximum_temperature_c: float
    minimum_raw_moisture: float
    clipped_cell_count: int
    maximum_clip_magnitude: float
    radial_monotonicity_all_outputs: bool
    non_center_argmax_count: int
    maximum_normalized_mass_balance_residual: float


@dataclass(frozen=True)
class SimulationResult:
    time_s: np.ndarray
    radius_cm: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    argmax_radius_cm: np.ndarray
    max_moisture: np.ndarray
    center_moisture: np.ndarray
    radial_monotonicity: np.ndarray
    drying_event_time_s: float
    report_end_time_s: int
    event_temperature_c: np.ndarray
    event_moisture: np.ndarray
    event_internal_radius_cm: np.ndarray
    event_internal_moisture: np.ndarray
    diagnostics: SolverDiagnostics


def drying_event_value(state: np.ndarray, node_count: int, threshold: float) -> float:
    return float(np.max(state[node_count:]) - threshold)


def _boundary_values(
    t_s: float,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
) -> tuple[float, float]:
    if t_s <= environment.terminal_time_s:
        return environment.interpolated_temperature(t_s), environment.interpolated_moisture(t_s)
    return post_boundary.temperature_c, post_boundary.moisture


def build_coupled_rhs(
    geometry: RadialGeometry,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
    config: SolverConfig,
) -> Callable[[float, np.ndarray], np.ndarray]:
    node_count = geometry.node_count

    def rhs(t_s: float, state: np.ndarray) -> np.ndarray:
        temperature = state[:node_count]
        moisture = state[node_count:]
        node_k = thermal_conductivity(moisture)
        node_d = moisture_diffusivity(moisture, temperature)
        face_k = interface_values(node_k, config.interface_mean)
        face_d = interface_values(node_d, config.interface_mean)
        heat_numerator = internal_flux_numerator(temperature, face_k, geometry)
        moisture_numerator = internal_flux_numerator(moisture, face_d, geometry)
        external_temperature, external_moisture = _boundary_values(t_s, environment, post_boundary)
        heat_numerator[-1] += (
            config.heat_transfer_w_m2_k
            * geometry.surface_area_m2_per_m
            * (external_temperature - temperature[-1])
        )
        moisture_numerator[-1] += (
            config.mass_transfer_m_s
            * geometry.surface_area_m2_per_m
            * (external_moisture - moisture[-1])
        )
        rho_cp = density(moisture) * heat_capacity(moisture)
        temperature_rate = heat_numerator / (rho_cp * geometry.volumes_m3_per_m)
        moisture_rate = moisture_numerator / geometry.volumes_m3_per_m
        return np.concatenate([temperature_rate, moisture_rate])

    return rhs


def _evaluate_segments(segments: list, times: np.ndarray) -> np.ndarray:
    if times.size == 0:
        raise ValueError("times must not be empty")
    state_count = int(segments[0].y.shape[0])
    evaluated = np.empty((state_count, times.size), dtype=float)
    assigned = np.zeros(times.size, dtype=bool)
    for segment in segments:
        mask = (~assigned) & (times >= segment.t[0] - 1e-8) & (times <= segment.t[-1] + 1e-8)
        if np.any(mask):
            evaluated[:, mask] = segment.sol(times[mask])
            assigned[mask] = True
    if not assigned.all():
        raise RuntimeError(f"no dense-output segment covers {times[~assigned]}")
    return evaluated


def _normalized_mass_balance_residual(
    t_s: float,
    state: np.ndarray,
    rhs: Callable[[float, np.ndarray], np.ndarray],
    geometry: RadialGeometry,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
    config: SolverConfig,
) -> float:
    n = geometry.node_count
    moisture = state[n:]
    rate = rhs(t_s, state)[n:]
    _, external_moisture = _boundary_values(t_s, environment, post_boundary)
    surface_flux = (
        config.mass_transfer_m_s
        * geometry.surface_area_m2_per_m
        * (external_moisture - moisture[-1])
    )
    residual = float(np.sum(rate * geometry.volumes_m3_per_m) - surface_flux)
    scale = float(np.sum(np.abs(rate * geometry.volumes_m3_per_m)) + abs(surface_flux) + 1e-30)
    return abs(residual) / scale


def solve_problem3_bdf(
    environment: EnvironmentSeries,
    config: SolverConfig | None = None,
    post_boundary: PostBoundary | None = None,
) -> SimulationResult:
    """两阶段积分耦合状态 [T_0..T_N,C_0..C_N]，以全域最大水分触发终止事件。"""

    config = config or SolverConfig()
    post_boundary = post_boundary or post_boundary_from_tail(environment)
    if environment.time_s[0] > 0.0:
        raise ValueError("environment must start at or before t=0")
    if config.initial_time_limit_s <= environment.terminal_time_s:
        raise ValueError("initial_time_limit_s must exceed the attachment terminal time")
    if config.fallback_time_limit_s <= config.initial_time_limit_s:
        raise ValueError("fallback_time_limit_s must exceed initial_time_limit_s")
    if config.extended_time_limit_s is not None and config.extended_time_limit_s <= config.fallback_time_limit_s:
        raise ValueError("extended_time_limit_s must exceed fallback_time_limit_s")
    if config.output_interval_s <= 0:
        raise ValueError("output_interval_s must be positive")

    geometry = build_radial_geometry(config.radius_cm, config.internal_dr_cm)
    internal_radius_cm = geometry.radius_m * 100.0
    output_geometry = build_radial_geometry(config.radius_cm, config.output_dr_cm)
    output_radius_cm = output_geometry.radius_m * 100.0
    node_count = geometry.node_count
    jacobian_pattern = build_coupled_jacobian_sparsity(node_count)
    rhs = build_coupled_rhs(geometry, environment, post_boundary, config)
    initial_state = np.concatenate(
        [
            np.full(node_count, config.initial_temperature_c, dtype=float),
            np.full(node_count, config.initial_moisture, dtype=float),
        ]
    )
    atol = np.concatenate(
        [
            np.full(node_count, config.temperature_absolute_tolerance),
            np.full(node_count, config.moisture_absolute_tolerance),
        ]
    )

    def integrate(start: float, end: float, state: np.ndarray, with_event: bool):
        events = None
        if with_event:
            def event(t_s: float, y: np.ndarray) -> float:
                del t_s
                return drying_event_value(y, node_count, config.completion_threshold)

            event.terminal = True
            event.direction = -1
            events = event
        solution = solve_ivp(
            rhs,
            (float(start), float(end)),
            state,
            method="BDF",
            events=events,
            dense_output=True,
            rtol=config.relative_tolerance,
            atol=atol,
            max_step=config.max_time_step_s,
            jac_sparsity=jacobian_pattern,
        )
        if not solution.success:
            raise RuntimeError(f"BDF integration failed: {solution.message}")
        if not np.isfinite(solution.y).all():
            raise RuntimeError("BDF integration returned non-finite values")
        return solution

    segments = []
    phase1 = integrate(0.0, environment.terminal_time_s, initial_state, False)
    segments.append(phase1)
    phase1_state = phase1.y[:, -1]
    if drying_event_value(phase1_state, node_count, config.completion_threshold) <= 0.0:
        raise RuntimeError("drying threshold was reached before the attachment boundary ended")

    phase2 = integrate(environment.terminal_time_s, config.initial_time_limit_s, phase1_state, True)
    segments.append(phase2)
    first_limit_reached = len(phase2.t_events[0]) == 0
    fallback_limit_reached = False
    event_time = None
    event_state = None
    if not first_limit_reached:
        event_time = float(phase2.t_events[0][0])
        event_state = phase2.y_events[0][0]
    else:
        phase3 = integrate(config.initial_time_limit_s, config.fallback_time_limit_s, phase2.y[:, -1], True)
        segments.append(phase3)
        if len(phase3.t_events[0]) == 0:
            fallback_limit_reached = True
            if config.extended_time_limit_s is None:
                terminal_max = float(np.max(phase3.y[node_count:, -1]))
                terminal_argmax = int(np.argmax(phase3.y[node_count:, -1]))
                raise RuntimeError(
                    f"drying threshold was not reached by {config.fallback_time_limit_s / 3600.0:.1f} h; "
                    f"max moisture={terminal_max:.9g} at r={internal_radius_cm[terminal_argmax]:.9g} cm"
                )
            phase4 = integrate(
                config.fallback_time_limit_s,
                config.extended_time_limit_s,
                phase3.y[:, -1],
                True,
            )
            segments.append(phase4)
            if len(phase4.t_events[0]) == 0:
                terminal_max = float(np.max(phase4.y[node_count:, -1]))
                terminal_argmax = int(np.argmax(phase4.y[node_count:, -1]))
                raise RuntimeError(
                    f"drying threshold was not reached by {config.extended_time_limit_s / 3600.0:.1f} h; "
                    f"max moisture={terminal_max:.9g} at r={internal_radius_cm[terminal_argmax]:.9g} cm"
                )
            event_time = float(phase4.t_events[0][0])
            event_state = phase4.y_events[0][0]
        else:
            event_time = float(phase3.t_events[0][0])
            event_state = phase3.y_events[0][0]

    assert event_time is not None and event_state is not None
    report_end = int(config.output_interval_s * ceil(event_time / config.output_interval_s))
    if report_end <= event_time + 1e-9:
        report_end += config.output_interval_s
    extension = integrate(event_time, float(report_end), event_state, False)
    while drying_event_value(extension.y[:, -1], node_count, config.completion_threshold) >= 0.0:
        report_end += config.output_interval_s
        extension = integrate(event_time, float(report_end), event_state, False)
    segments.append(extension)

    output_times = np.arange(config.output_interval_s, report_end + 1, config.output_interval_s, dtype=float)
    internal_states = _evaluate_segments(segments, output_times)
    internal_temperature = internal_states[:node_count, :].T
    internal_moisture = internal_states[node_count:, :].T
    minimum_raw_moisture = float(min(np.min(internal_moisture), np.min(event_state[node_count:])))
    if minimum_raw_moisture < config.material_negative_tolerance:
        raise RuntimeError(f"materially negative moisture returned: {minimum_raw_moisture:.6g}")

    output_temperature = sample_profiles(internal_radius_cm, output_radius_cm, internal_temperature)
    output_moisture = sample_profiles(internal_radius_cm, output_radius_cm, internal_moisture)
    event_temperature = sample_profiles(
        internal_radius_cm,
        output_radius_cm,
        event_state[:node_count][None, :],
    )[0]
    event_moisture = sample_profiles(
        internal_radius_cm,
        output_radius_cm,
        event_state[node_count:][None, :],
    )[0]
    argmax_indices = np.argmax(internal_moisture, axis=1)
    argmax_radius = internal_radius_cm[argmax_indices]
    max_moisture = np.max(internal_moisture, axis=1)
    monotonicity = np.all(np.diff(internal_moisture, axis=1) <= config.monotonicity_tolerance, axis=1)
    mass_residuals = np.array(
        [
            _normalized_mass_balance_residual(
                float(t_s), internal_states[:, index], rhs, geometry, environment, post_boundary, config
            )
            for index, t_s in enumerate(output_times)
        ]
    )
    all_solutions = segments
    diagnostics = SolverDiagnostics(
        internal_node_count=node_count,
        state_size=2 * node_count,
        function_evaluations=int(sum(solution.nfev for solution in all_solutions)),
        jacobian_evaluations=int(sum(solution.njev for solution in all_solutions)),
        lu_decompositions=int(sum(solution.nlu for solution in all_solutions)),
        phase1_end_time_s=environment.terminal_time_s,
        first_limit_reached_without_event=first_limit_reached,
        fallback_limit_reached_without_event=fallback_limit_reached,
        event_count=1,
        event_direction=-1,
        minimum_temperature_c=float(min(np.min(internal_temperature), np.min(event_state[:node_count]))),
        maximum_temperature_c=float(max(np.max(internal_temperature), np.max(event_state[:node_count]))),
        minimum_raw_moisture=minimum_raw_moisture,
        clipped_cell_count=0,
        maximum_clip_magnitude=0.0,
        radial_monotonicity_all_outputs=bool(monotonicity.all()),
        non_center_argmax_count=int(np.count_nonzero(argmax_indices)),
        maximum_normalized_mass_balance_residual=float(np.max(mass_residuals)),
    )
    return SimulationResult(
        time_s=output_times,
        radius_cm=output_radius_cm,
        temperature_c=output_temperature,
        moisture=output_moisture,
        argmax_radius_cm=argmax_radius,
        max_moisture=max_moisture,
        center_moisture=internal_moisture[:, 0],
        radial_monotonicity=monotonicity,
        drying_event_time_s=event_time,
        report_end_time_s=report_end,
        event_temperature_c=event_temperature,
        event_moisture=event_moisture,
        event_internal_radius_cm=internal_radius_cm,
        event_internal_moisture=event_state[node_count:].copy(),
        diagnostics=diagnostics,
    )



# ============================== 求解核心 · 后向欧拉-Picard 对照求解器 ==============================
@dataclass(frozen=True)
class BESolverConfig:
    radius_cm: float = 2.0
    internal_dr_cm: float = 0.025
    output_dr_cm: float = 0.1
    time_step_s: float = 10.0
    output_interval_s: int = 60
    initial_temperature_c: float = 28.0
    initial_moisture: float = 2.55
    completion_threshold: float = 0.15
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8e-7
    interface_mean: str = "harmonic"
    maximum_picard_iterations: int = 20
    temperature_picard_tolerance: float = 1e-8
    moisture_picard_tolerance: float = 1e-10
    initial_time_limit_s: float = 72.0 * 3600.0
    fallback_time_limit_s: float = 96.0 * 3600.0
    event_bracket_tolerance_s: float = 0.5
    material_negative_tolerance: float = -1e-8


@dataclass(frozen=True)
class BEDiagnostics:
    accepted_steps: int
    bisection_steps: int
    maximum_picard_iterations_used: int
    mean_picard_iterations: float
    unconverged_steps: int
    maximum_final_temperature_iteration_error: float
    maximum_final_moisture_iteration_error: float
    minimum_raw_moisture: float


@dataclass(frozen=True)
class BEResult:
    time_s: np.ndarray
    radius_cm: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    drying_event_time_s: float
    event_temperature_c: np.ndarray
    event_moisture: np.ndarray
    event_internal_radius_cm: np.ndarray
    event_internal_moisture: np.ndarray
    diagnostics: BEDiagnostics


def _solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    n = diagonal.size
    if n < 2:
        raise ValueError("tridiagonal system must have at least two rows")
    _, _, _, result, info = dgtsv(
        lower.copy(),
        diagonal.copy(),
        upper.copy(),
        right_hand_side.copy(),
        overwrite_dl=True,
        overwrite_d=True,
        overwrite_du=True,
        overwrite_b=True,
    )
    if info != 0:
        raise RuntimeError(f"compiled TDMA failed with LAPACK info={info}")
    return np.asarray(result, dtype=float)


def _implicit_fvm_step(
    previous: np.ndarray,
    node_transport: np.ndarray,
    capacity: np.ndarray,
    dt_s: float,
    boundary_transfer: float,
    boundary_external: float,
    geometry: RadialGeometry,
    interface_mean: str,
) -> np.ndarray:
    face_transport = interface_values(node_transport, interface_mean)
    conductance = face_transport * geometry.interface_areas_m2_per_m / geometry.dr_m
    accumulation = capacity * geometry.volumes_m3_per_m / dt_s
    lower = -conductance.copy()
    upper = -conductance.copy()
    diagonal = accumulation.copy()
    diagonal[:-1] += conductance
    diagonal[1:] += conductance
    right_hand_side = accumulation * previous
    boundary_conductance = boundary_transfer * geometry.surface_area_m2_per_m
    diagonal[-1] += boundary_conductance
    right_hand_side[-1] += boundary_conductance * boundary_external
    return _solve_tridiagonal(lower, diagonal, upper, right_hand_side)


def _boundary_at(t_s: float, environment: EnvironmentSeries, post_boundary: PostBoundary) -> tuple[float, float]:
    if t_s <= environment.terminal_time_s:
        return environment.interpolated_temperature(t_s), environment.interpolated_moisture(t_s)
    return post_boundary.temperature_c, post_boundary.moisture


def coupled_backward_euler_step(
    previous_temperature: np.ndarray,
    previous_moisture: np.ndarray,
    t_next_s: float,
    dt_s: float,
    geometry: RadialGeometry,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
    config: BESolverConfig,
) -> tuple[np.ndarray, np.ndarray, int, float, float]:
    if dt_s <= 0.0:
        raise ValueError("dt_s must be positive")
    external_temperature, external_moisture = _boundary_at(t_next_s, environment, post_boundary)
    temperature_guess = previous_temperature.copy()
    moisture_guess = previous_moisture.copy()
    last_temperature_error = np.inf
    last_moisture_error = np.inf
    for iteration in range(1, config.maximum_picard_iterations + 1):
        rho_cp = density(moisture_guess) * heat_capacity(moisture_guess)
        next_temperature = _implicit_fvm_step(
            previous_temperature,
            thermal_conductivity(moisture_guess),
            rho_cp,
            dt_s,
            config.heat_transfer_w_m2_k,
            external_temperature,
            geometry,
            config.interface_mean,
        )
        next_moisture = _implicit_fvm_step(
            previous_moisture,
            moisture_diffusivity(moisture_guess, next_temperature),
            np.ones_like(moisture_guess),
            dt_s,
            config.mass_transfer_m_s,
            external_moisture,
            geometry,
            config.interface_mean,
        )
        last_temperature_error = float(np.max(np.abs(next_temperature - temperature_guess)))
        last_moisture_error = float(np.max(np.abs(next_moisture - moisture_guess)))
        if (
            last_temperature_error <= config.temperature_picard_tolerance
            and last_moisture_error <= config.moisture_picard_tolerance
        ):
            return next_temperature, next_moisture, iteration, last_temperature_error, last_moisture_error
        temperature_guess = next_temperature
        moisture_guess = next_moisture
    raise RuntimeError(
        "Picard iteration did not converge: "
        f"t={t_next_s:.6f}s, dt={dt_s:.6f}s, "
        f"temperature_error={last_temperature_error:.3e}, moisture_error={last_moisture_error:.3e}"
    )


def solve_problem3_be(
    environment: EnvironmentSeries,
    config: BESolverConfig | None = None,
    post_boundary: PostBoundary | None = None,
) -> BEResult:
    config = config or BESolverConfig()
    post_boundary = post_boundary or post_boundary_from_tail(environment)
    if environment.time_s[0] > 0.0:
        raise ValueError("environment must start at or before t=0")
    geometry = build_radial_geometry(config.radius_cm, config.internal_dr_cm)
    internal_radius_cm = geometry.radius_m * 100.0
    output_radius_cm = build_radial_geometry(config.radius_cm, config.output_dr_cm).radius_m * 100.0
    temperature = np.full(geometry.node_count, config.initial_temperature_c, dtype=float)
    moisture = np.full(geometry.node_count, config.initial_moisture, dtype=float)
    time_s = 0.0
    output_times: list[float] = []
    output_temperature: list[np.ndarray] = []
    output_moisture: list[np.ndarray] = []
    iteration_counts: list[int] = []
    final_temperature_errors: list[float] = []
    final_moisture_errors: list[float] = []
    minimum_raw_moisture = float(np.min(moisture))
    bisection_steps = 0
    event_time = None
    event_temperature = None
    event_moisture = None

    while time_s < config.fallback_time_limit_s - 1e-12:
        step_end = min(time_s + config.time_step_s, config.fallback_time_limit_s)
        if time_s < environment.terminal_time_s < step_end:
            step_end = environment.terminal_time_s
        next_output = config.output_interval_s * (np.floor(time_s / config.output_interval_s) + 1.0)
        if time_s + 1e-10 < next_output < step_end - 1e-10:
            step_end = float(next_output)
        dt_s = step_end - time_s
        next_temperature, next_moisture, iterations, error_t, error_c = coupled_backward_euler_step(
            temperature,
            moisture,
            step_end,
            dt_s,
            geometry,
            environment,
            post_boundary,
            config,
        )
        iteration_counts.append(iterations)
        final_temperature_errors.append(error_t)
        final_moisture_errors.append(error_c)
        minimum_raw_moisture = min(minimum_raw_moisture, float(np.min(next_moisture)))
        if minimum_raw_moisture < config.material_negative_tolerance:
            raise RuntimeError(f"materially negative moisture returned: {minimum_raw_moisture:.6g}")

        previous_event_value = float(np.max(moisture) - config.completion_threshold)
        next_event_value = float(np.max(next_moisture) - config.completion_threshold)
        if previous_event_value > 0.0 and next_event_value <= 0.0:
            left_time = time_s
            left_temperature = temperature
            left_moisture = moisture
            right_time = step_end
            right_temperature = next_temperature
            right_moisture = next_moisture
            while right_time - left_time > config.event_bracket_tolerance_s:
                middle_time = 0.5 * (left_time + right_time)
                middle_temperature, middle_moisture, _, _, _ = coupled_backward_euler_step(
                    left_temperature,
                    left_moisture,
                    middle_time,
                    middle_time - left_time,
                    geometry,
                    environment,
                    post_boundary,
                    config,
                )
                bisection_steps += 1
                if float(np.max(middle_moisture) - config.completion_threshold) <= 0.0:
                    right_time = middle_time
                    right_temperature = middle_temperature
                    right_moisture = middle_moisture
                else:
                    left_time = middle_time
                    left_temperature = middle_temperature
                    left_moisture = middle_moisture
            event_time = right_time
            event_temperature = right_temperature
            event_moisture = right_moisture
            break

        time_s = step_end
        temperature = next_temperature
        moisture = next_moisture
        if np.isclose(time_s % config.output_interval_s, 0.0, atol=1e-8):
            output_times.append(time_s)
            output_temperature.append(temperature.copy())
            output_moisture.append(moisture.copy())

    if event_time is None or event_temperature is None or event_moisture is None:
        raise RuntimeError(
            f"drying threshold was not reached by {config.fallback_time_limit_s / 3600.0:.1f} h"
        )
    if not output_times:
        raise RuntimeError("no regular output was produced")

    output_temperature_internal = np.vstack(output_temperature)
    output_moisture_internal = np.vstack(output_moisture)
    diagnostics = BEDiagnostics(
        accepted_steps=len(iteration_counts),
        bisection_steps=bisection_steps,
        maximum_picard_iterations_used=int(max(iteration_counts)),
        mean_picard_iterations=float(np.mean(iteration_counts)),
        unconverged_steps=0,
        maximum_final_temperature_iteration_error=float(max(final_temperature_errors)),
        maximum_final_moisture_iteration_error=float(max(final_moisture_errors)),
        minimum_raw_moisture=minimum_raw_moisture,
    )
    return BEResult(
        time_s=np.asarray(output_times),
        radius_cm=output_radius_cm,
        temperature_c=sample_profiles(internal_radius_cm, output_radius_cm, output_temperature_internal),
        moisture=sample_profiles(internal_radius_cm, output_radius_cm, output_moisture_internal),
        drying_event_time_s=float(event_time),
        event_temperature_c=sample_profiles(
            internal_radius_cm, output_radius_cm, event_temperature[None, :]
        )[0],
        event_moisture=sample_profiles(
            internal_radius_cm, output_radius_cm, event_moisture[None, :]
        )[0],
        event_internal_radius_cm=internal_radius_cm,
        event_internal_moisture=event_moisture.copy(),
        diagnostics=diagnostics,
    )



# ============================== 运行与产物（全部默认，无命令行参数）==============================
BASE = Path(__file__).resolve().parent                       # 代码同根目录
ATTACHMENT1 = BASE / "附件" / "附件1.xlsx"                    # 环境数据（输入）
TEMPLATE = BASE / "附件" / "附件3" / "result3.xlsx"           # 结果格式模板（题目附件3，单 sheet）
OUT_XLSX = BASE / "result3.xlsx"                             # 输出：同根目录
OUT_JSON = BASE / "verification3.json"
PAPER_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0])


def load_environment(path: Path) -> EnvironmentSeries:
    """pandas 读附件1 三列（时间/温度/水分浓度）→ EnvironmentSeries。"""

    frame = pd.read_excel(path)
    required = ["时间", "温度", "水分浓度"]
    missing = [name for name in required if name not in frame.columns]
    if missing:
        raise ValueError(f"附件1缺少列：{missing}")
    return EnvironmentSeries(
        frame["时间"].to_numpy(float),
        frame["温度"].to_numpy(float),
        frame["水分浓度"].to_numpy(float),
    )


def values_at_times(result: SimulationResult | BEResult, times_s: np.ndarray) -> np.ndarray:
    indices = []
    for value in times_s:
        matches = np.where(np.isclose(result.time_s, value, rtol=0.0, atol=1e-8))[0]
        if not len(matches):
            raise ValueError(f"result has no regular output at t={value}")
        indices.append(int(matches[0]))
    radius_indices = [int(np.where(np.isclose(result.radius_cm, radius))[0][0]) for radius in PAPER_RADII_CM]
    return result.moisture[np.ix_(indices, radius_indices)]


def common_paper_times(*results: SimulationResult | BEResult) -> np.ndarray:
    latest_common = min(float(result.drying_event_time_s) for result in results)
    return np.arange(6.0 * 3600.0, np.floor(latest_common / (6.0 * 3600.0)) * 6.0 * 3600.0 + 1.0, 6.0 * 3600.0)


def comparison(left: SimulationResult | BEResult, right: SimulationResult | BEResult) -> dict[str, object]:
    times = common_paper_times(left, right)
    left_values = values_at_times(left, times)
    right_values = values_at_times(right, times)
    difference = np.abs(left_values - right_values)
    location = np.unravel_index(int(np.argmax(difference)), difference.shape)
    return {
        "event_time_difference_s": float(abs(left.drying_event_time_s - right.drying_event_time_s)),
        "event_time_relative_difference": float(
            abs(left.drying_event_time_s - right.drying_event_time_s) / right.drying_event_time_s
        ),
        "paper_sample_maximum_moisture_difference": float(difference[location]),
        "paper_sample_time_h_at_maximum": float(times[location[0]] / 3600.0),
        "paper_sample_radius_cm_at_maximum": float(PAPER_RADII_CM[location[1]]),
    }


def summarize_bdf(result: SimulationResult, runtime_s: float) -> dict[str, object]:
    return {
        "event_time_s": result.drying_event_time_s,
        "event_time_h": result.drying_event_time_s / 3600.0,
        "t60_s": result.report_end_time_s,
        "event_max_moisture": float(result.event_internal_moisture.max()),
        "event_argmax_radius_cm": float(
            result.event_internal_radius_cm[int(np.argmax(result.event_internal_moisture))]
        ),
        "event_center_moisture": float(result.event_internal_moisture[0]),
        "event_surface_moisture": float(result.event_internal_moisture[-1]),
        "t60_max_moisture": float(result.max_moisture[-1]),
        "previous_60s_max_moisture": float(result.max_moisture[-2]) if len(result.max_moisture) > 1 else None,
        "runtime_s": runtime_s,
        "diagnostics": asdict(result.diagnostics),
    }


def summarize_be(result: BEResult, runtime_s: float) -> dict[str, object]:
    return {
        "event_time_s": result.drying_event_time_s,
        "event_time_h": result.drying_event_time_s / 3600.0,
        "event_max_moisture": float(result.event_internal_moisture.max()),
        "event_argmax_radius_cm": float(
            result.event_internal_radius_cm[int(np.argmax(result.event_internal_moisture))]
        ),
        "event_center_moisture": float(result.event_internal_moisture[0]),
        "event_surface_moisture": float(result.event_internal_moisture[-1]),
        "runtime_s": runtime_s,
        "diagnostics": asdict(result.diagnostics),
    }


def bdf_run(label: str, environment: EnvironmentSeries, config: SolverConfig, post: PostBoundary | None = None):
    print(f"START BDF {label}", flush=True)
    started = time.perf_counter()
    result = solve_problem3_bdf(environment, config, post)
    elapsed = time.perf_counter() - started
    print(f"DONE  BDF {label}: {result.drying_event_time_s / 3600.0:.9f} h ({elapsed:.2f} s)", flush=True)
    return result, elapsed


def be_run(label: str, environment: EnvironmentSeries, config: BESolverConfig, post: PostBoundary | None = None):
    print(f"START BE  {label}", flush=True)
    started = time.perf_counter()
    result = solve_problem3_be(environment, config, post)
    elapsed = time.perf_counter() - started
    print(f"DONE  BE  {label}: {result.drying_event_time_s / 3600.0:.9f} h ({elapsed:.2f} s)", flush=True)
    return result, elapsed


def write_result_xlsx(result: SimulationResult, template: Path, out_path: Path) -> None:
    """读附件3 模板（单 sheet）→ 写达标停机水分表：22 列（标签 + 21 个半径 0.0~2.0 cm）。

    A 列为整数秒（number_format "0"），B~V 列四位小数（"0.0000"），冻结首行，A 列宽 24。
    结果时间序列从 t=60s 起（求解器输出不含 t=0 初值行），逐行写到第 2 行起。
    """

    workbook = openpyxl.load_workbook(template)
    sheet = workbook.worksheets[0]
    radius = result.radius_cm
    sheet.cell(1, 1).value = "时间\\到药材中心的距离"
    for j in range(radius.size):
        sheet.cell(1, 2 + j).value = round(float(radius[j]), 10)
    for i in range(result.time_s.size):
        time_cell = sheet.cell(i + 2, 1)
        time_cell.value = int(round(result.time_s[i]))
        time_cell.number_format = "0"
        for j in range(radius.size):
            cell = sheet.cell(i + 2, 2 + j)
            cell.value = float(result.moisture[i, j])   # 存全精度，仅用 number_format 控制四位小数显示（与原 result3.xlsx 一致）
            cell.number_format = "0.0000"
    sheet.column_dimensions["A"].width = 24
    sheet.freeze_panes = "A2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def main() -> None:
    if not ATTACHMENT1.exists():
        raise SystemExit(
            f"未找到附件：{ATTACHMENT1}\n"
            "请把题目附件 附件1.xlsx 放到本脚本同根目录的 附件/ 文件夹下，再运行：python problem3.py"
        )
    if not TEMPLATE.exists():
        raise SystemExit(
            f"未找到结果格式模板：{TEMPLATE}\n"
            "请把题目附件3 的 result3.xlsx 放到 附件/附件3/ 下，再运行：python problem3.py"
        )
    environment = load_environment(ATTACHMENT1)
    bar = "=" * 64

    # ---- ① 空间网格逐级加密：观察收敛阶 p、Richardson 外推、调和平均-网格序列（论文模型检验章）----
    spatial_dr = [0.1, 0.05, 0.025, 0.0125, 0.00625, 0.003125, 0.0015625, 0.00078125]
    spatial_config = SolverConfig(
        relative_tolerance=1e-8,
        temperature_absolute_tolerance=1e-8,
        moisture_absolute_tolerance=1e-10,
        max_time_step_s=60.0,
        extended_time_limit_s=1000.0 * 3600.0,
    )
    spatial_results: dict[float, SimulationResult] = {}
    spatial_rows: list[dict] = []
    for dr in spatial_dr:
        result, elapsed = bdf_run(f"spatial dr={dr:g} cm", environment, replace(spatial_config, internal_dr_cm=dr))
        spatial_results[dr] = result
        spatial_rows.append({"dr_cm": dr, **summarize_bdf(result, elapsed)})

    spatial_pairs: list[dict] = []
    for coarse, fine in zip(spatial_dr[:-1], spatial_dr[1:], strict=True):
        metrics = comparison(spatial_results[coarse], spatial_results[fine])
        metrics.update(
            {
                "coarse_dr_cm": coarse,
                "fine_dr_cm": fine,
                "event_location_same": bool(
                    spatial_rows[spatial_dr.index(coarse)]["event_argmax_radius_cm"]
                    == spatial_rows[spatial_dr.index(fine)]["event_argmax_radius_cm"]
                ),
            }
        )
        metrics["accepted"] = bool(
            metrics["event_time_difference_s"] < 60.0
            and metrics["event_time_relative_difference"] < 1e-3
            and metrics["paper_sample_maximum_moisture_difference"] < 5e-5
            and metrics["event_location_same"]
        )
        spatial_pairs.append(metrics)

    selected_dr = spatial_dr[-1]
    if not spatial_pairs[-1]["accepted"]:
        raise RuntimeError(f"finest spatial pair did not pass: {spatial_pairs[-1]}")
    last_three_times = [spatial_results[dr].drying_event_time_s for dr in spatial_dr[-3:]]
    observed_order = float(
        np.log2(abs(last_three_times[0] - last_three_times[1]) / abs(last_three_times[1] - last_three_times[2]))
    )
    richardson_event_time_s = float(
        last_three_times[2]
        + (last_three_times[2] - last_three_times[1]) / (2.0**observed_order - 1.0)
    )

    # ---- ② BDF 时间容差三档（B1/B2/B3），主结果取 B2（论文模型检验章）----
    time_configs = {
        "B1": replace(
            SolverConfig(internal_dr_cm=selected_dr),
            relative_tolerance=1e-7,
            temperature_absolute_tolerance=1e-7,
            moisture_absolute_tolerance=1e-9,
            max_time_step_s=60.0,
        ),
        "B2": replace(
            SolverConfig(internal_dr_cm=selected_dr),
            relative_tolerance=1e-8,
            temperature_absolute_tolerance=1e-8,
            moisture_absolute_tolerance=1e-10,
            max_time_step_s=30.0,
        ),
        "B3": replace(
            SolverConfig(internal_dr_cm=selected_dr),
            relative_tolerance=1e-9,
            temperature_absolute_tolerance=1e-10,
            moisture_absolute_tolerance=1e-12,
            max_time_step_s=10.0,
        ),
    }
    bdf_time_results: dict[str, SimulationResult] = {}
    bdf_time_rows: list[dict] = []
    for name, config in time_configs.items():
        result, elapsed = bdf_run(f"time {name}", environment, config)
        bdf_time_results[name] = result
        bdf_time_rows.append({"name": name, "config": asdict(config), **summarize_bdf(result, elapsed)})
    bdf_time_comparisons = {
        "B1_vs_B2": comparison(bdf_time_results["B1"], bdf_time_results["B2"]),
        "B2_vs_B3": comparison(bdf_time_results["B2"], bdf_time_results["B3"]),
    }
    main_result = bdf_time_results["B2"]

    # ---- ③ 后向欧拉-Picard Δt=2.5s 交叉对照（与最严格 BDF B3 比，论文模型检验章）----
    be_config = BESolverConfig(internal_dr_cm=selected_dr, time_step_s=2.5)
    be_result, be_elapsed = be_run("time dt=2.5s", environment, be_config)
    cross_solver = comparison(bdf_time_results["B3"], be_result)

    # ---- ④ 4h 后边界敏感性：数据驱动延拓 / 工程 3×3 / 传质系数 ±10%（论文稳定性分析章）----
    default_post = post_boundary_from_tail(environment)
    data_posts = {
        "attachment_terminal_value": default_post,
        "last_30_min_time_average": post_boundary_from_tail(environment, 1800.0),
        "last_1_h_time_average": post_boundary_from_tail(environment, 3600.0),
    }
    boundary_cache: dict[tuple[float, float, float], tuple[SimulationResult, float]] = {}

    def boundary_result(post: PostBoundary, hm: float, label: str):
        key = (round(post.temperature_c, 12), round(post.moisture, 12), round(hm, 15))
        if key not in boundary_cache:
            if key == (
                round(default_post.temperature_c, 12),
                round(default_post.moisture, 12),
                round(time_configs["B2"].mass_transfer_m_s, 15),
            ):
                boundary_cache[key] = (
                    main_result,
                    next(row["runtime_s"] for row in bdf_time_rows if row["name"] == "B2"),
                )
            else:
                boundary_cache[key] = bdf_run(
                    f"boundary {label}",
                    environment,
                    replace(time_configs["B2"], mass_transfer_m_s=hm),
                    post,
                )
        return boundary_cache[key]

    data_rows: list[dict] = []
    for name, post in data_posts.items():
        result, elapsed = boundary_result(post, time_configs["B2"].mass_transfer_m_s, name)
        data_rows.append(
            {
                "scenario": name,
                "post_temperature_c": post.temperature_c,
                "post_moisture": post.moisture,
                "mass_transfer_m_s": time_configs["B2"].mass_transfer_m_s,
                **summarize_bdf(result, elapsed),
                "difference_from_default_s": result.drying_event_time_s - main_result.drying_event_time_s,
            }
        )

    engineering_rows: list[dict] = []
    for temperature_delta in [-2.0, 0.0, 2.0]:
        for moisture_factor in [0.9, 1.0, 1.1]:
            post = PostBoundary(
                default_post.temperature_c + temperature_delta,
                default_post.moisture * moisture_factor,
            )
            label = f"T{temperature_delta:+g}_C{moisture_factor:.1f}"
            result, elapsed = boundary_result(post, time_configs["B2"].mass_transfer_m_s, label)
            engineering_rows.append(
                {
                    "scenario": label,
                    "post_temperature_c": post.temperature_c,
                    "post_moisture": post.moisture,
                    "mass_transfer_m_s": time_configs["B2"].mass_transfer_m_s,
                    **summarize_bdf(result, elapsed),
                    "difference_from_default_s": result.drying_event_time_s - main_result.drying_event_time_s,
                }
            )

    hm_rows: list[dict] = []
    for factor in [0.9, 1.0, 1.1]:
        hm = time_configs["B2"].mass_transfer_m_s * factor
        label = f"hm_{factor:.1f}x"
        result, elapsed = boundary_result(default_post, hm, label)
        hm_rows.append(
            {
                "scenario": label,
                "post_temperature_c": default_post.temperature_c,
                "post_moisture": default_post.moisture,
                "mass_transfer_m_s": hm,
                **summarize_bdf(result, elapsed),
                "difference_from_default_s": result.drying_event_time_s - main_result.drying_event_time_s,
            }
        )

    # ---- ⑤ 数值验收判定 ----
    bdf_time_pass = bool(
        bdf_time_comparisons["B2_vs_B3"]["event_time_difference_s"] < 60.0
        and bdf_time_comparisons["B2_vs_B3"]["paper_sample_maximum_moisture_difference"] < 1e-4
    )
    cross_solver_pass = bool(
        cross_solver["event_time_difference_s"] < 60.0
        and cross_solver["paper_sample_maximum_moisture_difference"] < 1e-4
        and be_result.diagnostics.unconverged_steps == 0
    )
    quality_pass = bool(
        main_result.diagnostics.maximum_normalized_mass_balance_residual < 1e-6
        and main_result.diagnostics.minimum_raw_moisture >= -1e-8
        and main_result.diagnostics.radial_monotonicity_all_outputs
        and float(main_result.event_internal_radius_cm[np.argmax(main_result.event_internal_moisture)]) == 0.0
        and main_result.max_moisture[-1] < main_result.max_moisture[-2]
        and main_result.max_moisture[-1] < time_configs["B2"].completion_threshold
    )
    all_pass = bool(spatial_pairs[-1]["accepted"] and bdf_time_pass and cross_solver_pass and quality_pass)

    data_times = [row["event_time_s"] for row in data_rows]
    engineering_times = [row["event_time_s"] for row in engineering_rows]
    hm_times = [row["event_time_s"] for row in hm_rows]
    all_boundary_times = data_times + engineering_times + hm_times

    # ---- ⑥ 产物：verification3.json（结构化验证 + 画图数据源）+ result3.xlsx（按附件3 模板）----
    verification = {
        "source": {
            "attachment1": str(ATTACHMENT1),
            "environment_terminal_time_s": environment.terminal_time_s,
            "environment_terminal_temperature_c": environment.terminal_temperature_c,
            "environment_terminal_moisture": environment.terminal_moisture,
        },
        "spatial_grid": {
            "runs": spatial_rows,
            "adjacent_comparisons": spatial_pairs,
            "selected_internal_dr_cm": selected_dr,
            "selected_pair_passed": spatial_pairs[-1]["accepted"],
            "observed_order_from_last_three_event_times": observed_order,
            "richardson_extrapolated_event_time_s": richardson_event_time_s,
            "richardson_extrapolated_event_time_h": richardson_event_time_s / 3600.0,
            "selected_minus_richardson_s": spatial_results[selected_dr].drying_event_time_s - richardson_event_time_s,
        },
        "bdf_time_sensitivity": {
            "runs": bdf_time_rows,
            "comparisons": bdf_time_comparisons,
            "passed": bdf_time_pass,
        },
        "cross_solver": {
            "comparison": cross_solver,
            "be_summary": summarize_be(be_result, be_elapsed),
            "passed": cross_solver_pass,
        },
        "quality": {
            "passed": quality_pass,
            "diagnostics": asdict(main_result.diagnostics),
            "event_profile_internal": main_result.event_internal_moisture.tolist(),
            "event_radius_internal_cm": main_result.event_internal_radius_cm.tolist(),
        },
        "boundary_sensitivity": {
            "data_driven": data_rows,
            "engineering_3x3": engineering_rows,
            "mass_transfer": hm_rows,
            "data_driven_range_s": [float(min(data_times)), float(max(data_times))],
            "engineering_range_s": [float(min(engineering_times)), float(max(engineering_times))],
            "mass_transfer_range_s": [float(min(hm_times)), float(max(hm_times))],
            "all_scenarios_range_s": [float(min(all_boundary_times)), float(max(all_boundary_times))],
        },
        "final": {
            "all_numerical_acceptance_checks_passed": all_pass,
            "internal_dr_cm": selected_dr,
            "bdf_configuration": "B2",
            "continuous_event_time_s": main_result.drying_event_time_s,
            "continuous_event_time_h": main_result.drying_event_time_s / 3600.0,
            "first_strictly_compliant_60s_time_s": main_result.report_end_time_s,
            "first_strictly_compliant_60s_time_h": main_result.report_end_time_s / 3600.0,
            "event_center_moisture": float(main_result.event_internal_moisture[0]),
            "event_surface_moisture": float(main_result.event_internal_moisture[-1]),
            "t60_max_moisture": float(main_result.max_moisture[-1]),
            "previous_60s_max_moisture": float(main_result.max_moisture[-2]),
            "boundary_all_scenarios_range_h": [
                float(min(all_boundary_times) / 3600.0),
                float(max(all_boundary_times) / 3600.0),
            ],
        },
    }
    OUT_JSON.write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")
    write_result_xlsx(main_result, TEMPLATE, OUT_XLSX)

    # ---- ⑦ 终端分节打印关键结论 ----
    argmax_radius = float(main_result.event_internal_radius_cm[np.argmax(main_result.event_internal_moisture)])
    print(bar)
    print("【问题三 主答案】固定半径 R=2cm、附录3 变物性热湿耦合，全截面达标停机的连续临界时间")
    print(f"  t*  = {main_result.drying_event_time_s:.4f} s = {main_result.drying_event_time_s / 3600.0:.4f} h")
    print(f"  t60 = {main_result.report_end_time_s} s = {main_result.report_end_time_s / 3600.0:.4f} h（首个严格达标 60s 时刻）")
    print(f"  事件时刻：中心水分 {main_result.event_internal_moisture[0]:.6f}、表面 {main_result.event_internal_moisture[-1]:.6f}、argmax 半径 {argmax_radius:.4f} cm")

    print(bar)
    print("【网格收敛性】空间 8 级加密（调和平均界面通量），事件时间 t*：")
    for row in spatial_rows:
        print(f"  Δr={row['dr_cm']:<12g} t*={row['event_time_h']:.4f} h  ({row['runtime_s']:.1f} s)")
    finest = spatial_pairs[-1]
    print(f"  末两级事件时间差 = {finest['event_time_difference_s']:.2f} s，相对差 = {finest['event_time_relative_difference']:.3e}")
    print(f"  观察收敛阶 p = {observed_order:.2f}，Richardson 外推 = {richardson_event_time_s:.4f} s = {richardson_event_time_s / 3600.0:.4f} h")
    print(f"  生产网格比外推高 {spatial_results[selected_dr].drying_event_time_s - richardson_event_time_s:.2f} s")

    print(bar)
    print("【BDF 时间容差三档】（主结果取 B2）")
    print(f"  B1 vs B2 事件时间差 = {bdf_time_comparisons['B1_vs_B2']['event_time_difference_s']:.3e} s")
    print(f"  B2 vs B3 事件时间差 = {bdf_time_comparisons['B2_vs_B3']['event_time_difference_s']:.3e} s")

    print(bar)
    print("【双求解器交叉对照】最严格 BDF(B3) vs 后向欧拉-Picard Δt=2.5s")
    print(f"  事件时间差 = {cross_solver['event_time_difference_s']:.2f} s，论文采样最大水分差 = {cross_solver['paper_sample_maximum_moisture_difference']:.3e}")
    print(f"  BE 诊断：accepted={be_result.diagnostics.accepted_steps}、bisection={be_result.diagnostics.bisection_steps}、未收敛步={be_result.diagnostics.unconverged_steps}、最小原始水分={be_result.diagnostics.minimum_raw_moisture:.3e}")

    print(bar)
    print("【守恒与物理核验】")
    print(f"  归一化质量守恒残差 = {main_result.diagnostics.maximum_normalized_mass_balance_residual:.3e}")
    print(f"  最小原始水分 = {main_result.diagnostics.minimum_raw_moisture:.3e}，裁剪单元数 = {main_result.diagnostics.clipped_cell_count}（未触发裁剪 = {main_result.diagnostics.clipped_cell_count == 0}）")
    print(f"  径向单调（全部输出）= {main_result.diagnostics.radial_monotonicity_all_outputs}，非中心 argmax 计数 = {main_result.diagnostics.non_center_argmax_count}")
    print(f"  t60 全截面最大水分 {main_result.max_moisture[-1]:.6f} < 前一时刻 {main_result.max_moisture[-2]:.6f}，且 < 阈值 {time_configs['B2'].completion_threshold}")
    print(f"  质量核验通过 = {quality_pass}")

    print(bar)
    print("【4h 后边界敏感性】")
    print("  数据驱动延拓（末值 / 30min 均值 / 1h 均值）：")
    for row in data_rows:
        print(f"    {row['scenario']:<28} t*={row['event_time_h']:.4f} h（Δ={row['difference_from_default_s']:+.1f} s）")
    print(f"    区间 [{min(data_times) / 3600.0:.4f}, {max(data_times) / 3600.0:.4f}] h")
    print("  工程情景 3×3（温度 ±2℃ × 水分 ×0.9/1.0/1.1）：")
    for row in engineering_rows:
        print(f"    {row['scenario']:<12} t*={row['event_time_h']:.4f} h（Δ={row['difference_from_default_s']:+.1f} s）")
    print(f"    区间 [{min(engineering_times) / 3600.0:.4f}, {max(engineering_times) / 3600.0:.4f}] h")
    print("  传质系数 hm ±10%：")
    for row in hm_rows:
        print(f"    {row['scenario']:<12} t*={row['event_time_h']:.4f} h（Δ={row['difference_from_default_s']:+.1f} s）")
    print(f"    区间 [{min(hm_times) / 3600.0:.4f}, {max(hm_times) / 3600.0:.4f}] h")
    print(f"  全部情景 t* 区间 [{min(all_boundary_times) / 3600.0:.4f}, {max(all_boundary_times) / 3600.0:.4f}] h")

    print(bar)
    print(f"产物已写出：{OUT_XLSX.name}（达标停机水分表，{main_result.time_s.size} 行 × {main_result.radius_cm.size + 1} 列）、{OUT_JSON.name}")
    print(f"全部数值验收通过 = {all_pass}")
    print(bar)


if __name__ == "__main__":
    main()
