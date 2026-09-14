"""问题一：一维径向守恒有限体积 + 隐式 BDF，求 30 min 预热阶段的温度场与水分场。

相对原版两处改进（默认行为，无开关）：界面扩散系数谐波平均、输出前水分 C≥0 裁剪。
温度场与水分场在问题一解耦，改进只影响水分场。main 内自动做「网格无关性」与
「算术→谐波改进效果」两项对比（论文模型检验章引用：加密网格温度/水分最大变化、改进水分变化）。

直接运行（无需任何参数）：python problem1.py
  · 读同根目录下 附件/附件1.xlsx（环境数据）与 附件/附件3/result1.xlsx（结果格式模板），读不到给提示；
  · 按模板格式把结果写到同根目录 result1.xlsx，验证摘要写 verification1.json；
  · 终端分节打印表1(温度)/表2(水分) + 网格无关 + 改进效果 + 物理核验。
"""
from __future__ import annotations

import json
from dataclasses import asdict, dataclass, replace
from pathlib import Path

import numpy as np
import openpyxl
from scipy.integrate import solve_ivp
from scipy.sparse import diags


# ============================== 求解核心 ==============================
@dataclass(frozen=True)
class EnvironmentSeries:
    """Piecewise-linear oven boundary data."""

    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def __post_init__(self) -> None:
        sizes = {len(self.time_s), len(self.temperature_c), len(self.moisture)}
        if len(sizes) != 1 or not len(self.time_s):
            raise ValueError("environment arrays must have the same non-zero length")
        if np.any(np.diff(self.time_s) <= 0):
            raise ValueError("environment times must be strictly increasing")

    def temperature_at(self, t_s: float) -> float:
        return float(np.interp(t_s, self.time_s, self.temperature_c))

    def moisture_at(self, t_s: float) -> float:
        return float(np.interp(t_s, self.time_s, self.moisture))


@dataclass(frozen=True)
class SolverConfig:
    radius_cm: float = 2.0
    internal_dr_cm: float = 0.0025
    output_dr_cm: float = 0.1
    duration_s: int = 1800
    output_dt_s: float = 1.0
    initial_temperature_c: float = 28.0
    initial_moisture: float = 2.55
    density_kg_m3: float = 820.0
    heat_capacity_j_kg_k: float = 2600.0
    thermal_conductivity_w_m_k: float = 0.36
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8e-7
    relative_tolerance: float = 1e-7
    absolute_tolerance: float = 1e-9
    max_time_step_s: float = 10.0
    interface_mean: str = "harmonic"          # 改进①：界面 D 谐波平均
    clip_negative_moisture: bool = True       # 改进②：输出前 C≥0 裁剪


@dataclass(frozen=True)
class SimulationResult:
    time_s: np.ndarray
    radius_cm: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray


def build_radial_grid_cm(radius_cm: float, dr_cm: float) -> np.ndarray:
    count = int(round(radius_cm / dr_cm)) + 1
    if count < 2 or not np.isclose(radius_cm / dr_cm, count - 1):
        raise ValueError("radius_cm must be an integer multiple of dr_cm")
    return np.linspace(0.0, radius_cm, count)


def moisture_diffusivity(c: np.ndarray) -> np.ndarray:
    """Problem-1 diffusivity, D(C)=7e-9 exp(-0.89/C), in m²/s."""

    safe_c = np.maximum(np.asarray(c, dtype=float), 1e-12)
    return 7e-9 * np.exp(-0.89 / safe_c)


def interface_average(d_left: np.ndarray, d_right: np.ndarray, kind: str) -> np.ndarray:
    """控制体积界面扩散系数平均。

    kind="arithmetic"：0.5*(D_l+D_r)（原版）。
    kind="harmonic"：2*D_l*D_r/(D_l+D_r)（改进①）；对 D 跨数量级的串联扩散更准，
        且当任一侧 D→0（表面强脱水）时界面 D→0，比算术平均更物理。防 0/0：分母为 0 时取 0。
    """

    if kind == "arithmetic":
        return 0.5 * (d_left + d_right)
    if kind != "harmonic":
        raise ValueError(f"unknown interface_mean: {kind!r} (expected 'harmonic' or 'arithmetic')")
    denom = d_left + d_right
    safe = np.where(denom > 0.0, denom, 1.0)
    return np.where(denom > 0.0, 2.0 * d_left * d_right / safe, 0.0)


def _radial_geometry(radius_m: float, dr_m: float) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    node_count = int(round(radius_m / dr_m)) + 1
    radius = np.linspace(0.0, radius_m, node_count)
    left_faces = np.maximum(0.0, radius - dr_m / 2.0)
    right_faces = np.minimum(radius_m, radius + dr_m / 2.0)
    volumes = np.pi * (right_faces**2 - left_faces**2)
    interface_areas = 2.0 * np.pi * (radius[:-1] + dr_m / 2.0)
    return radius, volumes, interface_areas


def _add_internal_fluxes(
    derivative_numerator: np.ndarray,
    values: np.ndarray,
    interface_transport: np.ndarray,
    interface_areas: np.ndarray,
    dr_m: float,
) -> None:
    flux = interface_transport * interface_areas * (values[1:] - values[:-1]) / dr_m
    derivative_numerator[:-1] += flux
    derivative_numerator[1:] -= flux


def _sample_to_output_grid(
    internal_radius_cm: np.ndarray,
    output_radius_cm: np.ndarray,
    states_by_time: np.ndarray,
) -> np.ndarray:
    spacing_cm = internal_radius_cm[1] - internal_radius_cm[0]
    indices = np.rint(output_radius_cm / spacing_cm).astype(int)
    if np.allclose(internal_radius_cm[indices], output_radius_cm):
        return states_by_time[:, indices]
    return np.vstack(
        [np.interp(output_radius_cm, internal_radius_cm, row) for row in states_by_time]
    )


def solve_problem1(environment: EnvironmentSeries, config: SolverConfig | None = None) -> SimulationResult:
    """Solve the decoupled heat/moisture model by radial FVM + implicit BDF."""

    config = config or SolverConfig()
    internal_radius_cm = build_radial_grid_cm(config.radius_cm, config.internal_dr_cm)
    output_radius_cm = build_radial_grid_cm(config.radius_cm, config.output_dr_cm)
    dr_m = config.internal_dr_cm / 100.0
    radius_m = config.radius_cm / 100.0
    _, volumes, interface_areas = _radial_geometry(radius_m, dr_m)
    surface_area = 2.0 * np.pi * radius_m
    node_count = len(internal_radius_cm)
    jacobian_pattern = diags(
        [np.ones(node_count - 1), np.ones(node_count), np.ones(node_count - 1)],
        offsets=[-1, 0, 1],
        shape=(node_count, node_count),
        format="csc",
    )
    output_time_s = np.arange(
        0.0,
        config.duration_s + config.output_dt_s / 2.0,
        config.output_dt_s,
    )

    def temperature_rhs(t_s: float, temperature: np.ndarray) -> np.ndarray:
        numerator = np.zeros_like(temperature)
        _add_internal_fluxes(
            numerator,
            temperature,
            np.full(node_count - 1, config.thermal_conductivity_w_m_k),
            interface_areas,
            dr_m,
        )
        numerator[-1] += (
            config.heat_transfer_w_m2_k
            * surface_area
            * (environment.temperature_at(t_s) - temperature[-1])
        )
        return numerator / (config.density_kg_m3 * config.heat_capacity_j_kg_k * volumes)

    def moisture_rhs(t_s: float, moisture: np.ndarray) -> np.ndarray:
        numerator = np.zeros_like(moisture)
        node_diffusivity = moisture_diffusivity(moisture)
        # 改进①：界面 D 用谐波平均（config.interface_mean="arithmetic" 时退化为原版算术平均）
        interface_diffusivity = interface_average(
            node_diffusivity[:-1], node_diffusivity[1:], config.interface_mean
        )
        _add_internal_fluxes(
            numerator,
            moisture,
            interface_diffusivity,
            interface_areas,
            dr_m,
        )
        numerator[-1] += (
            config.mass_transfer_m_s
            * surface_area
            * (environment.moisture_at(t_s) - moisture[-1])
        )
        return numerator / volumes

    common_options = {
        "method": "BDF",
        "t_eval": output_time_s,
        "rtol": config.relative_tolerance,
        "atol": config.absolute_tolerance,
        "max_step": config.max_time_step_s,
        "jac_sparsity": jacobian_pattern,
    }
    temperature_solution = solve_ivp(
        temperature_rhs,
        (0.0, float(config.duration_s)),
        np.full(node_count, config.initial_temperature_c),
        **common_options,
    )
    moisture_solution = solve_ivp(
        moisture_rhs,
        (0.0, float(config.duration_s)),
        np.full(node_count, config.initial_moisture),
        **common_options,
    )
    if not temperature_solution.success:
        raise RuntimeError(f"temperature integration failed: {temperature_solution.message}")
    if not moisture_solution.success:
        raise RuntimeError(f"moisture integration failed: {moisture_solution.message}")

    temperature = _sample_to_output_grid(
        internal_radius_cm,
        output_radius_cm,
        temperature_solution.y.T,
    )
    moisture = _sample_to_output_grid(
        internal_radius_cm,
        output_radius_cm,
        moisture_solution.y.T,
    )
    # 改进②：表面强脱水可能出现非物理负水分，输出前裁剪到 C≥0（problem1 工况实测不触发，防御性）
    if config.clip_negative_moisture:
        moisture = np.maximum(moisture, 0.0)
    return SimulationResult(output_time_s, output_radius_cm, temperature, moisture)


# ============================== 运行与产物（全部默认，无命令行参数）==============================
BASE = Path(__file__).resolve().parent                       # 代码同根目录
ATTACHMENT1 = BASE / "附件" / "附件1.xlsx"                    # 环境数据（输入）
TEMPLATE = BASE / "附件" / "附件3" / "result1.xlsx"           # 结果格式模板（题目附件3）
OUT_XLSX = BASE / "result1.xlsx"                             # 输出：同根目录
OUT_JSON = BASE / "verification1.json"
SAMPLE_TIMES = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
SAMPLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0], dtype=float)
TEMP_SHEET = "温度"
MOIST_SHEET = "水分浓度"


def load_environment(path: Path) -> EnvironmentSeries:
    """openpyxl 读附件1 三列（时间/温度/水分浓度）→ EnvironmentSeries。"""

    worksheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = [r for r in worksheet.iter_rows(min_row=2, values_only=True) if r[0] not in (None, "")]
    return EnvironmentSeries(
        time_s=np.asarray([float(r[0]) for r in rows], dtype=float),
        temperature_c=np.asarray([float(r[1]) for r in rows], dtype=float),
        moisture=np.asarray([float(r[2]) for r in rows], dtype=float),
    )


def selected_values(result: SimulationResult, values: np.ndarray) -> np.ndarray:
    time_indices = [int(np.where(np.isclose(result.time_s, value))[0][0]) for value in SAMPLE_TIMES]
    radius_indices = [int(np.where(np.isclose(result.radius_cm, value))[0][0]) for value in SAMPLE_RADII_CM]
    return values[np.ix_(time_indices, radius_indices)]


def difference_summary(left: np.ndarray, right: np.ndarray) -> dict[str, float]:
    """两场逐点对照：最大绝对差 + 平均绝对差（论文模型检验章引用）。"""

    absolute = np.abs(left - right)
    return {
        "maximum_absolute_difference": float(absolute.max()),
        "mean_absolute_difference": float(absolute.mean()),
    }


def write_result_xlsx(result: SimulationResult, template: Path, out_path: Path) -> None:
    """读附件3 格式模板 → 按完整时间×半径网格覆盖写温度/水分两表（全精度存储，四位小数显示）。"""

    workbook = openpyxl.load_workbook(template)
    radius = result.radius_cm
    for sheet_name, field in ((TEMP_SHEET, result.temperature_c), (MOIST_SHEET, result.moisture)):
        sheet = workbook[sheet_name]
        sheet.cell(1, 1).value = "时间\\到药材中心的距离"
        for j in range(radius.size):
            sheet.cell(1, 2 + j).value = round(float(radius[j]), 10)
        for i in range(1, result.time_s.size):  # t=1..1800（跳过 t=0 初值行）
            sheet.cell(i + 1, 1).value = int(round(result.time_s[i]))
            for j in range(radius.size):
                cell = sheet.cell(i + 1, 2 + j)
                cell.value = float(field[i, j])  # 全精度存储，仅用 number_format 控制四位小数显示（与原 result1.xlsx 一致）
                cell.number_format = "0.0000"
        sheet.freeze_panes = "B2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def main() -> None:
    if not ATTACHMENT1.exists():
        raise SystemExit(
            f"未找到附件：{ATTACHMENT1}\n"
            "请把题目附件 附件1.xlsx 放到本脚本同根目录的 附件/ 文件夹下，再运行：python problem1.py"
        )
    if not TEMPLATE.exists():
        raise SystemExit(
            f"未找到结果格式模板：{TEMPLATE}\n"
            "请把题目附件3 的 result1.xlsx 放到 附件/附件3/ 下，再运行：python problem1.py"
        )
    environment = load_environment(ATTACHMENT1)

    improved_config = SolverConfig()                                                   # 默认：谐波平均 + C≥0 裁剪（主结果）
    baseline_config = SolverConfig(interface_mean="arithmetic", clip_negative_moisture=False)  # 原版行为（对照）
    grid_config = replace(improved_config, internal_dr_cm=improved_config.internal_dr_cm / 2.0)  # 网格加密一倍

    print("求解 baseline（算术平均、无裁剪，对照）...", flush=True)
    baseline = solve_problem1(environment, baseline_config)
    print("求解 improved（谐波平均、C≥0 裁剪，主结果）...", flush=True)
    improved = solve_problem1(environment, improved_config)
    print(f"求解 grid-refined（内部网格 {grid_config.internal_dr_cm} cm）...", flush=True)
    grid = solve_problem1(environment, grid_config)

    grid_independence = {
        "note": f"内部网格 {improved_config.internal_dr_cm} vs {grid_config.internal_dr_cm} cm（同一输出网格）",
        "temperature_c": difference_summary(improved.temperature_c, grid.temperature_c),
        "moisture": difference_summary(improved.moisture, grid.moisture),
    }
    improvement_effect = {
        "note": "谐波平均 vs 算术平均（同环境、同网格，唯一变量=界面平均方式+裁剪）；温度场解耦故逐点一致",
        "temperature_identical": bool(np.array_equal(improved.temperature_c, baseline.temperature_c)),
        "moisture": difference_summary(improved.moisture, baseline.moisture),
    }
    physical_checks = {
        "moisture_min": float(improved.moisture.min()),
        "moisture_max": float(improved.moisture.max()),
        "temperature_min_c": float(improved.temperature_c.min()),
        "temperature_max_c": float(improved.temperature_c.max()),
        "clip_is_noop": bool(improved.moisture.min() > 0.0),
        "all_finite": bool(np.isfinite(improved.temperature_c).all() and np.isfinite(improved.moisture).all()),
    }
    verification = {
        "source_attachment": str(ATTACHMENT1),
        "config": asdict(improved_config),
        "grid_independence": grid_independence,
        "improvement_effect": improvement_effect,
        "physical_checks": physical_checks,
        "selected_tables": {
            "temperature_c": [[float(t), *[float(v) for v in row]] for t, row in zip(SAMPLE_TIMES, selected_values(improved, improved.temperature_c))],
            "moisture": [[float(t), *[float(v) for v in row]] for t, row in zip(SAMPLE_TIMES, selected_values(improved, improved.moisture))],
        },
    }

    write_result_xlsx(improved, TEMPLATE, OUT_XLSX)
    OUT_JSON.write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")

    bar = "=" * 64
    radii_header = "  ".join(f"{r:>7.1f}" for r in SAMPLE_RADII_CM)
    print(bar, flush=True)
    print("问题一 30 min 预热 —— 温度/水分径向分布（谐波平均界面 D + C≥0 裁剪）", flush=True)
    print(bar, flush=True)
    print("[表1] 温度/℃（行=时间/s，列=距中心半径/cm）", flush=True)
    print(f"{'t/s':>6}  {radii_header}", flush=True)
    for t, row in zip(SAMPLE_TIMES, selected_values(improved, improved.temperature_c)):
        print(f"{int(t):>6}  " + "  ".join(f"{v:>7.4f}" for v in row), flush=True)
    print("[表2] 水分浓度/(kg/kg)", flush=True)
    print(f"{'t/s':>6}  {radii_header}", flush=True)
    for t, row in zip(SAMPLE_TIMES, selected_values(improved, improved.moisture)):
        print(f"{int(t):>6}  " + "  ".join(f"{v:>7.4f}" for v in row), flush=True)
    print("-" * 64, flush=True)
    print(f"[网格无关性] 内部网格减半后：温度最大变化={grid_independence['temperature_c']['maximum_absolute_difference']:.3e} ℃，"
          f"水分最大变化={grid_independence['moisture']['maximum_absolute_difference']:.3e} kg/kg", flush=True)
    print(f"[改进效果] 算术→谐波平均：温度场逐点一致={improvement_effect['temperature_identical']}，"
          f"水分最大变化={improvement_effect['moisture']['maximum_absolute_difference']:.3e} kg/kg", flush=True)
    print(f"[物理核验] 水分 min={physical_checks['moisture_min']:.6f} max={physical_checks['moisture_max']:.6f}；"
          f"温度 min={physical_checks['temperature_min_c']:.4f} max={physical_checks['temperature_max_c']:.4f} ℃；"
          f"C≥0 裁剪为 no-op={physical_checks['clip_is_noop']}", flush=True)
    print(bar, flush=True)
    print(f"已写出：{OUT_XLSX}（温度/水分两表，按附件3 模板格式）、{OUT_JSON}", flush=True)


if __name__ == "__main__":
    main()
