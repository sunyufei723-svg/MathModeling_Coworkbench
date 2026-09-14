"""问题四：附件2半径收缩 + 附录4物性下的热风烘干达标停机时间（自包含单文件）。

直接运行（无命令行参数，自动读取同根目录附件）：
    python problem4.py

输入附件（与本文件位于同一根目录）：
    附件/附件1.xlsx         烘房环境边界（时间 / 温度 / 水分浓度）
    附件/附件2.xlsx         药材半径收缩轨迹（时间 / 半径）
    附件/附件3/result4.xlsx  题目给定的结果输出格式模板

输出（写到与本文件同一根目录）：
    result4.xlsx        全截面水分浓度剖面（按附件3模板格式，全精度存储、四位小数显示）
    verification4.json  结构化验证结果（网格收敛 / 容差 / 双求解器 / 边界情景 / 收缩 / 坐标形式，兼画图数据源）
    终端分节打印        主答案 t*、各项验证结论与验收判定

模型：物质坐标 front-fixing 守恒径向有限体积（FVM）+ 联合状态自适应 BDF，
界面热导率 / 扩散系数取调和平均，连续事件 g(t)=max(C)-0.15 定位全截面达标停机时刻。
主答案 t* = 50.8245 h（N=3201 生产网格）。
"""

from __future__ import annotations

import json
import time
from dataclasses import asdict, dataclass, field, replace
from math import floor
from pathlib import Path
from typing import Callable

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter
from scipy.integrate import solve_ivp
from scipy.interpolate import PchipInterpolator
from scipy.linalg.lapack import dgtsv
from scipy.sparse import bmat, csc_matrix, diags


# ====== 求解核心 · 物性与环境边界（附录4）======


@dataclass(frozen=True)
class PhysicalParameters:
    """附录4物性及收缩幅度的可注入参数；默认值严格保持原模型口径。"""

    density_intercept: float = 760.0
    density_moisture_coefficient: float = 90.0
    diffusivity_prefactor: float = 4.2e-4
    diffusivity_moisture_parameter: float = 0.30
    activation_temperature_k: float = 3850.0
    shrinkage_amplitude_factor: float = 1.0

    def __post_init__(self) -> None:
        values = np.asarray(
            [
                self.density_intercept,
                self.density_moisture_coefficient,
                self.diffusivity_prefactor,
                self.diffusivity_moisture_parameter,
                self.activation_temperature_k,
                self.shrinkage_amplitude_factor,
            ],
            dtype=float,
        )
        if not np.isfinite(values).all() or np.any(values <= 0.0):
            raise ValueError("physical parameters must be positive and finite")


def safe_moisture(c: np.ndarray | float) -> np.ndarray:
    """只在物性公式中使用的数值保护，不修改求解状态。"""

    return np.maximum(np.asarray(c, dtype=float), 1.0e-9)


def density(
    c: np.ndarray | float,
    parameters: PhysicalParameters | None = None,
) -> np.ndarray:
    parameters = parameters or PhysicalParameters()
    c_safe = safe_moisture(c)
    return parameters.density_intercept + parameters.density_moisture_coefficient * c_safe


def heat_capacity(c: np.ndarray | float) -> np.ndarray:
    c_safe = safe_moisture(c)
    return 1850.0 + 2150.0 * c_safe / (1.0 + c_safe)


def thermal_conductivity(c: np.ndarray | float) -> np.ndarray:
    c_safe = safe_moisture(c)
    return 0.12 + 0.20 * c_safe / (1.0 + c_safe)


def moisture_diffusivity(
    c: np.ndarray | float,
    temperature_c: np.ndarray | float,
    parameters: PhysicalParameters | None = None,
) -> np.ndarray:
    parameters = parameters or PhysicalParameters()
    c_safe = safe_moisture(c)
    temperature_k = np.asarray(temperature_c, dtype=float) + 273.15
    if np.any(temperature_k <= 0.0):
        raise ValueError("temperature must be above absolute zero")
    return (
        parameters.diffusivity_prefactor
        * np.exp(-parameters.diffusivity_moisture_parameter / c_safe)
        * np.exp(-parameters.activation_temperature_k / temperature_k)
    )


@dataclass(frozen=True)
class EnvironmentSeries:
    """附件1边界序列；附件区间内仅作分段线性插值。"""

    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def __post_init__(self) -> None:
        time = np.asarray(self.time_s, dtype=float)
        temperature = np.asarray(self.temperature_c, dtype=float)
        moisture = np.asarray(self.moisture, dtype=float)
        if not (time.ndim == temperature.ndim == moisture.ndim == 1):
            raise ValueError("environment arrays must be one-dimensional")
        if len(time) < 2 or not (len(time) == len(temperature) == len(moisture)):
            raise ValueError("environment arrays must have equal length >= 2")
        if np.any(np.diff(time) <= 0.0):
            raise ValueError("environment time must be strictly increasing")
        if not (np.isfinite(time).all() and np.isfinite(temperature).all() and np.isfinite(moisture).all()):
            raise ValueError("environment arrays must be finite")
        object.__setattr__(self, "time_s", time)
        object.__setattr__(self, "temperature_c", temperature)
        object.__setattr__(self, "moisture", moisture)

    @property
    def terminal_time_s(self) -> float:
        return float(self.time_s[-1])

    @property
    def terminal_temperature_c(self) -> float:
        return float(self.temperature_c[-1])

    @property
    def terminal_moisture(self) -> float:
        return float(self.moisture[-1])

    def values_at(self, t_s: float) -> tuple[float, float]:
        if t_s < self.time_s[0] or t_s > self.time_s[-1]:
            raise ValueError("interpolation time is outside the attachment interval")
        return (
            float(np.interp(t_s, self.time_s, self.temperature_c)),
            float(np.interp(t_s, self.time_s, self.moisture)),
        )


@dataclass(frozen=True)
class PostBoundary:
    temperature_c: float
    moisture: float


def post_boundary_from_tail(
    environment: EnvironmentSeries,
    averaging_window_s: float | None = None,
) -> PostBoundary:
    if averaging_window_s is None:
        return PostBoundary(environment.terminal_temperature_c, environment.terminal_moisture)
    if averaging_window_s <= 0.0:
        raise ValueError("averaging_window_s must be positive")
    start = max(float(environment.time_s[0]), environment.terminal_time_s - averaging_window_s)
    interior = environment.time_s[(environment.time_s > start) & (environment.time_s <= environment.terminal_time_s)]
    times = np.concatenate([[start], interior])
    temperature = np.interp(times, environment.time_s, environment.temperature_c)
    moisture = np.interp(times, environment.time_s, environment.moisture)
    duration = float(times[-1] - times[0])
    if duration <= 0.0:
        raise ValueError("averaging window has zero duration")
    return PostBoundary(
        float(np.trapezoid(temperature, times) / duration),
        float(np.trapezoid(moisture, times) / duration),
    )


def boundary_at(
    t_s: float,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
) -> tuple[float, float]:
    if t_s <= environment.terminal_time_s:
        return environment.values_at(t_s)
    return post_boundary.temperature_c, post_boundary.moisture


@dataclass(frozen=True)
class RadiusSeries:
    """附件2半径序列，默认线性插值，可切换单调 PCHIP。"""

    time_s: np.ndarray
    radius_cm: np.ndarray
    interpolation: str = "linear"
    _pchip: PchipInterpolator | None = field(default=None, init=False, repr=False, compare=False)

    def __post_init__(self) -> None:
        time = np.asarray(self.time_s, dtype=float)
        radius = np.asarray(self.radius_cm, dtype=float)
        if time.ndim != 1 or radius.ndim != 1 or len(time) != len(radius) or len(time) < 2:
            raise ValueError("radius arrays must be one-dimensional with equal length >= 2")
        if np.any(np.diff(time) <= 0.0):
            raise ValueError("radius time must be strictly increasing")
        if np.any(radius <= 0.0) or not (np.isfinite(time).all() and np.isfinite(radius).all()):
            raise ValueError("radius data must be positive and finite")
        if self.interpolation not in {"linear", "pchip"}:
            raise ValueError("interpolation must be 'linear' or 'pchip'")
        object.__setattr__(self, "time_s", time)
        object.__setattr__(self, "radius_cm", radius)
        if self.interpolation == "pchip":
            object.__setattr__(self, "_pchip", PchipInterpolator(time, radius, extrapolate=False))

    @property
    def terminal_time_s(self) -> float:
        return float(self.time_s[-1])

    def radius_cm_at(self, t_s: float) -> float:
        if t_s <= self.time_s[0]:
            return float(self.radius_cm[0])
        if t_s >= self.time_s[-1]:
            return float(self.radius_cm[-1])
        if self.interpolation == "linear":
            return float(np.interp(t_s, self.time_s, self.radius_cm))
        assert self._pchip is not None
        return float(self._pchip(t_s))

    def radius_rate_cm_s_at(self, t_s: float) -> float:
        if t_s <= self.time_s[0] or t_s >= self.time_s[-1]:
            return 0.0
        if self.interpolation == "linear":
            index = int(np.searchsorted(self.time_s, t_s, side="right") - 1)
            index = min(max(index, 0), len(self.time_s) - 2)
            return float(
                (self.radius_cm[index + 1] - self.radius_cm[index])
                / (self.time_s[index + 1] - self.time_s[index])
            )
        assert self._pchip is not None
        return float(self._pchip.derivative()(t_s))

    def with_interpolation(self, interpolation: str) -> "RadiusSeries":
        return RadiusSeries(self.time_s, self.radius_cm, interpolation)


@dataclass(frozen=True)
class BoundaryScenario:
    name: str
    post_boundary: PostBoundary
    mass_transfer_factor: float = 1.0


def engineering_boundary_scenarios(default: PostBoundary) -> list[BoundaryScenario]:
    return [
        BoundaryScenario(
            f"T{temperature_delta:+g}_C{moisture_factor:.1f}",
            PostBoundary(
                default.temperature_c + temperature_delta,
                default.moisture * moisture_factor,
            ),
        )
        for temperature_delta in (-2.0, 0.0, 2.0)
        for moisture_factor in (0.9, 1.0, 1.1)
    ]


def mass_transfer_scenarios(default: PostBoundary) -> list[BoundaryScenario]:
    return [BoundaryScenario(f"hm_{factor:.1f}x", default, factor) for factor in (0.9, 1.0, 1.1)]


# ====== 求解核心 · 物质坐标 FVM 几何与界面离散 ======


@dataclass(frozen=True)
class ReferenceGeometry:
    xi: np.ndarray
    volumes_hat: np.ndarray
    interface_areas_hat: np.ndarray
    dxi: float

    @property
    def node_count(self) -> int:
        return int(self.xi.size)


def build_reference_geometry(node_count: int) -> ReferenceGeometry:
    if node_count < 2:
        raise ValueError("node_count must be at least two")
    xi = np.linspace(0.0, 1.0, node_count)
    dxi = 1.0 / (node_count - 1)
    left_faces = np.maximum(0.0, xi - 0.5 * dxi)
    right_faces = np.minimum(1.0, xi + 0.5 * dxi)
    volumes_hat = 0.5 * (right_faces**2 - left_faces**2)
    interface_areas_hat = xi[:-1] + 0.5 * dxi
    return ReferenceGeometry(xi, volumes_hat, interface_areas_hat, dxi)


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    left = np.asarray(left, dtype=float)
    right = np.asarray(right, dtype=float)
    denominator = left + right
    result = np.zeros(np.broadcast_shapes(left.shape, right.shape), dtype=float)
    np.divide(2.0 * left * right, denominator, out=result, where=np.abs(denominator) > 0.0)
    return result


def interface_values(node_values: np.ndarray, mean: str) -> np.ndarray:
    values = np.asarray(node_values, dtype=float)
    if mean == "harmonic":
        return harmonic_mean(values[:-1], values[1:])
    if mean == "arithmetic":
        return 0.5 * (values[:-1] + values[1:])
    raise ValueError("mean must be 'harmonic' or 'arithmetic'")


def internal_flux_numerator(
    values: np.ndarray,
    interface_transport: np.ndarray,
    geometry: ReferenceGeometry,
) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    flux = (
        np.asarray(interface_transport, dtype=float)
        * geometry.interface_areas_hat
        * (values[1:] - values[:-1])
        / geometry.dxi
    )
    numerator = np.zeros_like(values)
    numerator[:-1] += flux
    numerator[1:] -= flux
    return numerator


def build_coupled_jacobian_sparsity(node_count: int) -> csc_matrix:
    if node_count < 2:
        raise ValueError("node_count must be at least two")
    tri = diags(
        [np.ones(node_count - 1), np.ones(node_count), np.ones(node_count - 1)],
        [-1, 0, 1],
        shape=(node_count, node_count),
        format="csc",
    )
    return bmat([[tri, tri], [tri, tri]], format="csc")


def sample_physical_profile(
    profile: np.ndarray,
    geometry: ReferenceGeometry,
    radius_cm: float,
    fixed_distances_cm: np.ndarray,
) -> np.ndarray:
    output = np.full(len(fixed_distances_cm) + 1, np.nan, dtype=float)
    for index, distance_cm in enumerate(fixed_distances_cm):
        if distance_cm <= radius_cm + 1.0e-12:
            output[index] = float(np.interp(distance_cm / radius_cm, geometry.xi, profile))
    output[-1] = float(profile[-1])
    return output



# ====== 求解核心 · 联合状态自适应 BDF 求解器（主求解器）======


DEFAULT_FIXED_DISTANCES_CM = tuple(float(value) for value in np.round(np.arange(0.0, 2.0, 0.1), 10))


@dataclass(frozen=True)
class SolverConfig:
    node_count: int = 801
    sample_interval_s: int = 60
    report_interval_s: int = 60
    fixed_distances_cm: tuple[float, ...] = DEFAULT_FIXED_DISTANCES_CM
    initial_temperature_c: float = 28.0
    initial_moisture: float = 2.55
    completion_threshold: float = 0.15
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8.0e-7
    k_interface_mean: str = "harmonic"
    d_interface_mean: str = "harmonic"
    model_form: str = "material"
    fixed_radius_cm: float | None = None
    physical_parameters: PhysicalParameters = field(default_factory=PhysicalParameters)
    integration_method: str = "BDF"
    relative_tolerance: float = 1.0e-8
    temperature_absolute_tolerance: float = 1.0e-8
    moisture_absolute_tolerance: float = 1.0e-10
    max_time_step_s: float = 30.0
    initial_time_limit_s: float = 72.0 * 3600.0
    fallback_time_limit_s: float = 96.0 * 3600.0
    material_negative_tolerance: float = -1.0e-8
    monotonicity_tolerance: float = 1.0e-8


@dataclass(frozen=True)
class SolverDiagnostics:
    internal_node_count: int
    state_size: int
    function_evaluations: int
    jacobian_evaluations: int
    lu_decompositions: int
    phase1_end_time_s: float
    first_limit_reached_without_event: bool
    event_count: int
    event_direction: int
    minimum_temperature_c: float
    maximum_temperature_c: float
    minimum_raw_moisture: float
    clipped_cell_count: int
    maximum_clip_magnitude: float
    radial_monotonicity_all_outputs: bool
    non_center_argmax_count: int
    maximum_normalized_balance_residual: float
    dry_mass_index_relative_range: float


@dataclass(frozen=True)
class SimulationResult:
    time_s: np.ndarray
    radius_cm: np.ndarray
    fixed_distance_cm: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    argmax_xi: np.ndarray
    argmax_radius_cm: np.ndarray
    max_moisture: np.ndarray
    center_moisture: np.ndarray
    radial_monotonicity: np.ndarray
    dry_mass_index: np.ndarray
    drying_event_time_s: float
    report_end_time_s: int
    previous_report_time_s: int
    previous_report_max_moisture: float
    report_max_moisture: float
    event_radius_cm: float
    event_temperature_internal: np.ndarray
    event_moisture_internal: np.ndarray
    event_output_temperature: np.ndarray
    event_output_moisture: np.ndarray
    diagnostics: SolverDiagnostics


def drying_event_value(state: np.ndarray, node_count: int, threshold: float) -> float:
    return float(np.max(state[node_count:]) - threshold)


def _radius_values(t_s: float, radius: RadiusSeries, config: SolverConfig) -> tuple[float, float]:
    if config.fixed_radius_cm is not None:
        if config.fixed_radius_cm <= 0.0:
            raise ValueError("fixed_radius_cm must be positive")
        return config.fixed_radius_cm / 100.0, 0.0
    base_radius_cm = float(radius.radius_cm[0])
    shrinkage_factor = config.physical_parameters.shrinkage_amplitude_factor
    radius_cm = base_radius_cm + shrinkage_factor * (radius.radius_cm_at(t_s) - base_radius_cm)
    radius_rate_cm_s = shrinkage_factor * radius.radius_rate_cm_s_at(t_s)
    if radius_cm <= 0.0:
        raise ValueError("scaled radius must remain positive")
    return radius_cm / 100.0, radius_rate_cm_s / max(radius_cm, 1.0e-30)


def build_coupled_rhs(
    geometry: ReferenceGeometry,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
    radius: RadiusSeries,
    config: SolverConfig,
) -> Callable[[float, np.ndarray], np.ndarray]:
    if config.model_form not in {"material", "eulerian"}:
        raise ValueError("model_form must be 'material' or 'eulerian'")
    node_count = geometry.node_count

    def rhs(t_s: float, state: np.ndarray) -> np.ndarray:
        temperature = state[:node_count]
        moisture = state[node_count:]
        radius_m, radius_rate_ratio_s = _radius_values(t_s, radius, config)
        external_temperature, external_moisture = boundary_at(t_s, environment, post_boundary)

        node_k = thermal_conductivity(moisture)
        node_d = moisture_diffusivity(moisture, temperature, config.physical_parameters)
        heat_numerator = internal_flux_numerator(
            temperature,
            interface_values(node_k, config.k_interface_mean),
            geometry,
        )
        moisture_numerator = internal_flux_numerator(
            moisture,
            interface_values(node_d, config.d_interface_mean),
            geometry,
        )
        heat_numerator[-1] += radius_m * config.heat_transfer_w_m2_k * (
            external_temperature - temperature[-1]
        )
        moisture_numerator[-1] += radius_m * config.mass_transfer_m_s * (
            external_moisture - moisture[-1]
        )

        rho_cp = density(moisture, config.physical_parameters) * heat_capacity(moisture)
        temperature_rate = heat_numerator / (radius_m**2 * rho_cp * geometry.volumes_hat)
        moisture_rate = moisture_numerator / (radius_m**2 * geometry.volumes_hat)

        if config.model_form == "eulerian":
            temperature_gradient = np.gradient(temperature, geometry.dxi, edge_order=2)
            moisture_gradient = np.gradient(moisture, geometry.dxi, edge_order=2)
            temperature_rate += geometry.xi * radius_rate_ratio_s * temperature_gradient
            moisture_rate += geometry.xi * radius_rate_ratio_s * moisture_gradient

        derivative = np.concatenate([temperature_rate, moisture_rate])
        if not np.isfinite(derivative).all():
            raise RuntimeError("RHS returned non-finite values")
        return derivative

    return rhs


def _evaluate_segments(segments: list, times: np.ndarray) -> np.ndarray:
    states = np.empty((segments[0].y.shape[0], times.size), dtype=float)
    assigned = np.zeros(times.size, dtype=bool)
    for segment in segments:
        mask = (~assigned) & (times >= segment.t[0] - 1.0e-8) & (times <= segment.t[-1] + 1.0e-8)
        if np.any(mask):
            states[:, mask] = segment.sol(times[mask])
            assigned[mask] = True
    if not assigned.all():
        raise RuntimeError(f"no dense-output segment covers {times[~assigned]}")
    return states


def _sample_times(event_time_s: float, report_end_time_s: int, config: SolverConfig) -> np.ndarray:
    regular = np.arange(
        config.sample_interval_s,
        floor(event_time_s / config.sample_interval_s) * config.sample_interval_s + 1,
        config.sample_interval_s,
        dtype=float,
    )
    required = np.array(
        [float(report_end_time_s - config.report_interval_s), float(report_end_time_s)],
        dtype=float,
    )
    return np.unique(np.concatenate([regular, required]))


def _normalized_balance_residual(
    t_s: float,
    state: np.ndarray,
    rhs: Callable[[float, np.ndarray], np.ndarray],
    geometry: ReferenceGeometry,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
    radius: RadiusSeries,
    config: SolverConfig,
) -> float:
    if config.model_form != "material":
        return 0.0
    n = geometry.node_count
    moisture = state[n:]
    radius_m, _ = _radius_values(t_s, radius, config)
    _, external_moisture = boundary_at(t_s, environment, post_boundary)
    boundary_term = config.mass_transfer_m_s * (external_moisture - moisture[-1]) / radius_m
    weighted_rate = float(np.sum(rhs(t_s, state)[n:] * geometry.volumes_hat))
    residual = weighted_rate - boundary_term
    scale = abs(weighted_rate) + abs(boundary_term) + 1.0e-30
    return abs(residual) / scale


def _dry_mass_index(radius_m: float, moisture: np.ndarray, geometry: ReferenceGeometry) -> float:
    return _dry_mass_index_with_parameters(radius_m, moisture, geometry, PhysicalParameters())


def _dry_mass_index_with_parameters(
    radius_m: float,
    moisture: np.ndarray,
    geometry: ReferenceGeometry,
    parameters: PhysicalParameters,
) -> float:
    return float(
        radius_m**2
        * np.sum(
            density(moisture, parameters)
            / (1.0 + np.asarray(moisture, dtype=float))
            * geometry.volumes_hat
        )
    )


def solve_problem4_bdf(
    environment: EnvironmentSeries,
    radius: RadiusSeries,
    config: SolverConfig | None = None,
    post_boundary: PostBoundary | None = None,
) -> SimulationResult:
    """严格物质坐标 FVM + 联合状态自适应 BDF + 连续终止事件。"""

    config = config or SolverConfig()
    post_boundary = post_boundary or post_boundary_from_tail(environment)
    if environment.time_s[0] > 0.0:
        raise ValueError("environment must start at or before t=0")
    if config.initial_time_limit_s <= environment.terminal_time_s:
        raise ValueError("initial_time_limit_s must exceed attachment terminal time")
    if config.fallback_time_limit_s <= config.initial_time_limit_s:
        raise ValueError("fallback_time_limit_s must exceed initial_time_limit_s")
    if config.sample_interval_s <= 0 or config.report_interval_s <= 0:
        raise ValueError("output intervals must be positive")
    if config.integration_method not in {"BDF", "Radau"}:
        raise ValueError("integration_method must be 'BDF' or 'Radau'")

    geometry = build_reference_geometry(config.node_count)
    node_count = geometry.node_count
    jacobian_pattern = build_coupled_jacobian_sparsity(node_count)
    rhs = build_coupled_rhs(geometry, environment, post_boundary, radius, config)
    initial_state = np.concatenate(
        [
            np.full(node_count, config.initial_temperature_c, dtype=float),
            np.full(node_count, config.initial_moisture, dtype=float),
        ]
    )
    atol = np.concatenate(
        [
            np.full(node_count, config.temperature_absolute_tolerance),
            np.full(node_count, config.moisture_absolute_tolerance),
        ]
    )

    def integrate(start: float, end: float, state: np.ndarray, with_event: bool):
        events = None
        if with_event:
            def event(_: float, y: np.ndarray) -> float:
                return drying_event_value(y, node_count, config.completion_threshold)

            event.terminal = True
            event.direction = -1
            events = event
        solution = solve_ivp(
            rhs,
            (float(start), float(end)),
            state,
            method=config.integration_method,
            events=events,
            dense_output=True,
            rtol=config.relative_tolerance,
            atol=atol,
            max_step=config.max_time_step_s,
            jac_sparsity=jacobian_pattern,
        )
        if not solution.success:
            raise RuntimeError(f"{config.integration_method} integration failed: {solution.message}")
        if not np.isfinite(solution.y).all():
            raise RuntimeError(f"{config.integration_method} integration returned non-finite values")
        return solution

    segments = []
    phase1 = integrate(0.0, environment.terminal_time_s, initial_state, False)
    segments.append(phase1)
    if drying_event_value(phase1.y[:, -1], node_count, config.completion_threshold) <= 0.0:
        raise RuntimeError("drying threshold was reached before attachment boundary ended")

    phase2 = integrate(
        environment.terminal_time_s,
        config.initial_time_limit_s,
        phase1.y[:, -1],
        True,
    )
    segments.append(phase2)
    first_limit_reached = len(phase2.t_events[0]) == 0
    if first_limit_reached:
        phase3 = integrate(
            config.initial_time_limit_s,
            config.fallback_time_limit_s,
            phase2.y[:, -1],
            True,
        )
        segments.append(phase3)
        if len(phase3.t_events[0]) == 0:
            terminal_max = float(np.max(phase3.y[node_count:, -1]))
            raise RuntimeError(
                f"drying threshold was not reached by {config.fallback_time_limit_s / 3600.0:.1f} h; "
                f"max moisture={terminal_max:.9g}"
            )
        event_time = float(phase3.t_events[0][0])
        event_state = phase3.y_events[0][0]
    else:
        event_time = float(phase2.t_events[0][0])
        event_state = phase2.y_events[0][0]

    report_end = int(config.report_interval_s * (floor(event_time / config.report_interval_s) + 1))
    extension = integrate(event_time, float(report_end), event_state, False)
    while drying_event_value(extension.y[:, -1], node_count, config.completion_threshold) >= 0.0:
        report_end += config.report_interval_s
        extension = integrate(event_time, float(report_end), event_state, False)
    segments.append(extension)

    previous_report = report_end - config.report_interval_s
    sample_times = _sample_times(event_time, report_end, config)
    sample_states = _evaluate_segments(segments, sample_times)
    sample_temperature = sample_states[:node_count, :].T
    sample_moisture = sample_states[node_count:, :].T
    radius_cm = np.array(
        [100.0 * _radius_values(float(t_s), radius, config)[0] for t_s in sample_times],
        dtype=float,
    )
    fixed_distance = np.asarray(config.fixed_distances_cm, dtype=float)
    output_temperature = np.vstack(
        [
            sample_physical_profile(profile, geometry, r_cm, fixed_distance)
            for profile, r_cm in zip(sample_temperature, radius_cm, strict=True)
        ]
    )
    output_moisture = np.vstack(
        [
            sample_physical_profile(profile, geometry, r_cm, fixed_distance)
            for profile, r_cm in zip(sample_moisture, radius_cm, strict=True)
        ]
    )

    argmax_indices = np.argmax(sample_moisture, axis=1)
    argmax_xi = geometry.xi[argmax_indices]
    argmax_radius_cm = argmax_xi * radius_cm
    max_moisture = np.max(sample_moisture, axis=1)
    monotonicity = np.all(np.diff(sample_moisture, axis=1) <= config.monotonicity_tolerance, axis=1)
    balance_residuals = np.array(
        [
            _normalized_balance_residual(
                float(t_s), sample_states[:, index], rhs, geometry, environment, post_boundary, radius, config
            )
            for index, t_s in enumerate(sample_times)
        ],
        dtype=float,
    )
    dry_mass = np.array(
        [
            _dry_mass_index_with_parameters(
                r_cm / 100.0,
                profile,
                geometry,
                config.physical_parameters,
            )
            for r_cm, profile in zip(radius_cm, sample_moisture, strict=True)
        ],
        dtype=float,
    )
    initial_radius_m, _ = _radius_values(0.0, radius, config)
    initial_dry_mass = _dry_mass_index_with_parameters(
        initial_radius_m,
        initial_state[node_count:],
        geometry,
        config.physical_parameters,
    )
    dry_mass_with_initial = np.concatenate([[initial_dry_mass], dry_mass])
    dry_mass_relative_range = float(
        (np.max(dry_mass_with_initial) - np.min(dry_mass_with_initial))
        / max(abs(initial_dry_mass), 1.0e-30)
    )

    event_radius_cm = 100.0 * _radius_values(event_time, radius, config)[0]
    event_temperature = event_state[:node_count].copy()
    event_moisture = event_state[node_count:].copy()
    event_output_temperature = sample_physical_profile(
        event_temperature, geometry, event_radius_cm, fixed_distance
    )
    event_output_moisture = sample_physical_profile(
        event_moisture, geometry, event_radius_cm, fixed_distance
    )
    previous_state = _evaluate_segments(segments, np.array([float(previous_report)]))[:, 0]
    report_state = extension.sol(float(report_end))
    minimum_raw_moisture = float(
        min(np.min(sample_moisture), np.min(event_moisture), np.min(report_state[node_count:]))
    )
    if minimum_raw_moisture < config.material_negative_tolerance:
        raise RuntimeError(f"materially negative moisture returned: {minimum_raw_moisture:.6g}")

    diagnostics = SolverDiagnostics(
        internal_node_count=node_count,
        state_size=2 * node_count,
        function_evaluations=int(sum(segment.nfev for segment in segments)),
        jacobian_evaluations=int(sum(segment.njev for segment in segments)),
        lu_decompositions=int(sum(segment.nlu for segment in segments)),
        phase1_end_time_s=environment.terminal_time_s,
        first_limit_reached_without_event=first_limit_reached,
        event_count=1,
        event_direction=-1,
        minimum_temperature_c=float(
            min(config.initial_temperature_c, np.min(sample_temperature), np.min(event_temperature))
        ),
        maximum_temperature_c=float(max(np.max(sample_temperature), np.max(event_temperature))),
        minimum_raw_moisture=minimum_raw_moisture,
        clipped_cell_count=0,
        maximum_clip_magnitude=0.0,
        radial_monotonicity_all_outputs=bool(monotonicity.all()),
        non_center_argmax_count=int(np.count_nonzero(argmax_indices)),
        maximum_normalized_balance_residual=float(np.max(balance_residuals)),
        dry_mass_index_relative_range=dry_mass_relative_range,
    )
    return SimulationResult(
        time_s=sample_times,
        radius_cm=radius_cm,
        fixed_distance_cm=fixed_distance,
        temperature_c=output_temperature,
        moisture=output_moisture,
        argmax_xi=argmax_xi,
        argmax_radius_cm=argmax_radius_cm,
        max_moisture=max_moisture,
        center_moisture=sample_moisture[:, 0],
        radial_monotonicity=monotonicity,
        dry_mass_index=dry_mass,
        drying_event_time_s=event_time,
        report_end_time_s=report_end,
        previous_report_time_s=previous_report,
        previous_report_max_moisture=float(np.max(previous_state[node_count:])),
        report_max_moisture=float(np.max(report_state[node_count:])),
        event_radius_cm=event_radius_cm,
        event_temperature_internal=event_temperature,
        event_moisture_internal=event_moisture,
        event_output_temperature=event_output_temperature,
        event_output_moisture=event_output_moisture,
        diagnostics=diagnostics,
    )



# ====== 求解核心 · 后向欧拉-Picard 对照求解器 ======


@dataclass(frozen=True)
class BESolverConfig:
    node_count: int = 801
    time_step_s: float = 10.0
    sample_interval_s: int = 21600
    fixed_distances_cm: tuple[float, ...] = DEFAULT_FIXED_DISTANCES_CM
    initial_temperature_c: float = 28.0
    initial_moisture: float = 2.55
    completion_threshold: float = 0.15
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8.0e-7
    k_interface_mean: str = "harmonic"
    d_interface_mean: str = "harmonic"
    fixed_radius_cm: float | None = None
    maximum_picard_iterations: int = 20
    temperature_picard_tolerance: float = 1.0e-8
    moisture_picard_tolerance: float = 1.0e-10
    fallback_time_limit_s: float = 96.0 * 3600.0
    event_bracket_tolerance_s: float = 0.5
    material_negative_tolerance: float = -1.0e-8


@dataclass(frozen=True)
class BEDiagnostics:
    accepted_steps: int
    bisection_steps: int
    maximum_picard_iterations_used: int
    mean_picard_iterations: float
    unconverged_steps: int
    maximum_final_temperature_iteration_error: float
    maximum_final_moisture_iteration_error: float
    minimum_raw_moisture: float


@dataclass(frozen=True)
class BEResult:
    time_s: np.ndarray
    radius_cm: np.ndarray
    fixed_distance_cm: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    drying_event_time_s: float
    event_radius_cm: float
    event_temperature_internal: np.ndarray
    event_moisture_internal: np.ndarray
    event_output_temperature: np.ndarray
    event_output_moisture: np.ndarray
    diagnostics: BEDiagnostics


def _solve_tridiagonal(
    lower: np.ndarray,
    diagonal: np.ndarray,
    upper: np.ndarray,
    right_hand_side: np.ndarray,
) -> np.ndarray:
    _, _, _, result, info = dgtsv(
        lower.copy(),
        diagonal.copy(),
        upper.copy(),
        right_hand_side.copy(),
        overwrite_dl=True,
        overwrite_d=True,
        overwrite_du=True,
        overwrite_b=True,
    )
    if info != 0:
        raise RuntimeError(f"compiled TDMA failed with LAPACK info={info}")
    return np.asarray(result, dtype=float)


def _radius_m_at(t_s: float, radius: RadiusSeries, fixed_radius_cm: float | None) -> float:
    if fixed_radius_cm is not None:
        if fixed_radius_cm <= 0.0:
            raise ValueError("fixed_radius_cm must be positive")
        return fixed_radius_cm / 100.0
    return radius.radius_cm_at(t_s) / 100.0


def _implicit_reference_step(
    previous: np.ndarray,
    node_transport: np.ndarray,
    capacity: np.ndarray,
    dt_s: float,
    radius_m: float,
    boundary_transfer: float,
    boundary_external: float,
    geometry: ReferenceGeometry,
    interface_mean: str,
) -> np.ndarray:
    face_transport = interface_values(node_transport, interface_mean)
    conductance = face_transport * geometry.interface_areas_hat / (geometry.dxi * radius_m**2)
    accumulation = capacity * geometry.volumes_hat / dt_s
    lower = -conductance.copy()
    upper = -conductance.copy()
    diagonal = accumulation.copy()
    diagonal[:-1] += conductance
    diagonal[1:] += conductance
    right_hand_side = accumulation * previous
    boundary_conductance = boundary_transfer / radius_m
    diagonal[-1] += boundary_conductance
    right_hand_side[-1] += boundary_conductance * boundary_external
    return _solve_tridiagonal(lower, diagonal, upper, right_hand_side)


def coupled_backward_euler_step(
    previous_temperature: np.ndarray,
    previous_moisture: np.ndarray,
    t_next_s: float,
    dt_s: float,
    geometry: ReferenceGeometry,
    environment: EnvironmentSeries,
    post_boundary: PostBoundary,
    radius: RadiusSeries,
    config: BESolverConfig,
) -> tuple[np.ndarray, np.ndarray, int, float, float]:
    if dt_s <= 0.0:
        raise ValueError("dt_s must be positive")
    external_temperature, external_moisture = boundary_at(t_next_s, environment, post_boundary)
    radius_m = _radius_m_at(t_next_s, radius, config.fixed_radius_cm)
    temperature_guess = previous_temperature.copy()
    moisture_guess = previous_moisture.copy()
    last_temperature_error = np.inf
    last_moisture_error = np.inf
    for iteration in range(1, config.maximum_picard_iterations + 1):
        rho_cp = density(moisture_guess) * heat_capacity(moisture_guess)
        next_temperature = _implicit_reference_step(
            previous_temperature,
            thermal_conductivity(moisture_guess),
            rho_cp,
            dt_s,
            radius_m,
            config.heat_transfer_w_m2_k,
            external_temperature,
            geometry,
            config.k_interface_mean,
        )
        next_moisture = _implicit_reference_step(
            previous_moisture,
            moisture_diffusivity(moisture_guess, next_temperature),
            np.ones_like(moisture_guess),
            dt_s,
            radius_m,
            config.mass_transfer_m_s,
            external_moisture,
            geometry,
            config.d_interface_mean,
        )
        last_temperature_error = float(np.max(np.abs(next_temperature - temperature_guess)))
        last_moisture_error = float(np.max(np.abs(next_moisture - moisture_guess)))
        if (
            last_temperature_error <= config.temperature_picard_tolerance
            and last_moisture_error <= config.moisture_picard_tolerance
        ):
            return next_temperature, next_moisture, iteration, last_temperature_error, last_moisture_error
        temperature_guess = next_temperature
        moisture_guess = next_moisture
    raise RuntimeError(
        "Picard iteration did not converge: "
        f"t={t_next_s:.6f}s, dt={dt_s:.6f}s, "
        f"temperature_error={last_temperature_error:.3e}, moisture_error={last_moisture_error:.3e}"
    )


def solve_problem4_be(
    environment: EnvironmentSeries,
    radius: RadiusSeries,
    config: BESolverConfig | None = None,
    post_boundary: PostBoundary | None = None,
) -> BEResult:
    config = config or BESolverConfig()
    post_boundary = post_boundary or post_boundary_from_tail(environment)
    geometry = build_reference_geometry(config.node_count)
    fixed_distance = np.asarray(config.fixed_distances_cm, dtype=float)
    temperature = np.full(config.node_count, config.initial_temperature_c, dtype=float)
    moisture = np.full(config.node_count, config.initial_moisture, dtype=float)
    time_s = 0.0
    output_times: list[float] = []
    output_radii: list[float] = []
    output_temperature: list[np.ndarray] = []
    output_moisture: list[np.ndarray] = []
    iteration_counts: list[int] = []
    temperature_errors: list[float] = []
    moisture_errors: list[float] = []
    minimum_raw_moisture = float(np.min(moisture))
    bisection_steps = 0
    event_time = None
    event_temperature = None
    event_moisture = None

    while time_s < config.fallback_time_limit_s - 1.0e-12:
        step_end = min(time_s + config.time_step_s, config.fallback_time_limit_s)
        if time_s < environment.terminal_time_s < step_end:
            step_end = environment.terminal_time_s
        next_sample = config.sample_interval_s * (np.floor(time_s / config.sample_interval_s) + 1.0)
        if time_s + 1.0e-10 < next_sample < step_end - 1.0e-10:
            step_end = float(next_sample)
        dt_s = step_end - time_s
        next_temperature, next_moisture, iterations, error_t, error_c = coupled_backward_euler_step(
            temperature,
            moisture,
            step_end,
            dt_s,
            geometry,
            environment,
            post_boundary,
            radius,
            config,
        )
        iteration_counts.append(iterations)
        temperature_errors.append(error_t)
        moisture_errors.append(error_c)
        minimum_raw_moisture = min(minimum_raw_moisture, float(np.min(next_moisture)))
        if minimum_raw_moisture < config.material_negative_tolerance:
            raise RuntimeError(f"materially negative moisture returned: {minimum_raw_moisture:.6g}")

        previous_event_value = float(np.max(moisture) - config.completion_threshold)
        next_event_value = float(np.max(next_moisture) - config.completion_threshold)
        if previous_event_value > 0.0 and next_event_value <= 0.0:
            left_time = time_s
            left_temperature = temperature.copy()
            left_moisture = moisture.copy()
            right_time = step_end
            right_temperature = next_temperature.copy()
            right_moisture = next_moisture.copy()
            while right_time - left_time > config.event_bracket_tolerance_s:
                middle_time = 0.5 * (left_time + right_time)
                middle_temperature, middle_moisture, _, _, _ = coupled_backward_euler_step(
                    left_temperature,
                    left_moisture,
                    middle_time,
                    middle_time - left_time,
                    geometry,
                    environment,
                    post_boundary,
                    radius,
                    config,
                )
                bisection_steps += 1
                if float(np.max(middle_moisture) - config.completion_threshold) <= 0.0:
                    right_time = middle_time
                    right_temperature = middle_temperature
                    right_moisture = middle_moisture
                else:
                    left_time = middle_time
                    left_temperature = middle_temperature
                    left_moisture = middle_moisture
            event_time = right_time
            event_temperature = right_temperature
            event_moisture = right_moisture
            break

        time_s = step_end
        temperature = next_temperature
        moisture = next_moisture
        if np.isclose(time_s % config.sample_interval_s, 0.0, atol=1.0e-8):
            radius_cm = 100.0 * _radius_m_at(time_s, radius, config.fixed_radius_cm)
            output_times.append(time_s)
            output_radii.append(radius_cm)
            output_temperature.append(sample_physical_profile(temperature, geometry, radius_cm, fixed_distance))
            output_moisture.append(sample_physical_profile(moisture, geometry, radius_cm, fixed_distance))

    if event_time is None or event_temperature is None or event_moisture is None:
        raise RuntimeError(
            f"drying threshold was not reached by {config.fallback_time_limit_s / 3600.0:.1f} h"
        )
    event_radius_cm = 100.0 * _radius_m_at(event_time, radius, config.fixed_radius_cm)
    diagnostics = BEDiagnostics(
        accepted_steps=len(iteration_counts),
        bisection_steps=bisection_steps,
        maximum_picard_iterations_used=int(max(iteration_counts)),
        mean_picard_iterations=float(np.mean(iteration_counts)),
        unconverged_steps=0,
        maximum_final_temperature_iteration_error=float(max(temperature_errors)),
        maximum_final_moisture_iteration_error=float(max(moisture_errors)),
        minimum_raw_moisture=minimum_raw_moisture,
    )
    return BEResult(
        time_s=np.asarray(output_times, dtype=float),
        radius_cm=np.asarray(output_radii, dtype=float),
        fixed_distance_cm=fixed_distance,
        temperature_c=np.vstack(output_temperature) if output_temperature else np.empty((0, len(fixed_distance) + 1)),
        moisture=np.vstack(output_moisture) if output_moisture else np.empty((0, len(fixed_distance) + 1)),
        drying_event_time_s=float(event_time),
        event_radius_cm=event_radius_cm,
        event_temperature_internal=event_temperature.copy(),
        event_moisture_internal=event_moisture.copy(),
        event_output_temperature=sample_physical_profile(
            event_temperature, geometry, event_radius_cm, fixed_distance
        ),
        event_output_moisture=sample_physical_profile(event_moisture, geometry, event_radius_cm, fixed_distance),
        diagnostics=diagnostics,
    )



# ====== 运行与产物（全部默认，无命令行参数）======


BASE = Path(__file__).resolve().parent
ATTACHMENT1 = BASE / "附件" / "附件1.xlsx"
ATTACHMENT2 = BASE / "附件" / "附件2.xlsx"
TEMPLATE = BASE / "附件" / "附件3" / "result4.xlsx"
OUT_XLSX = BASE / "result4.xlsx"
OUT_JSON = BASE / "verification4.json"


def json_dump(path: Path, value) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False), encoding="utf-8")


def load_inputs(attachment1: Path, attachment2: Path) -> tuple[EnvironmentSeries, RadiusSeries]:
    environment_frame = pd.read_excel(attachment1)
    radius_frame = pd.read_excel(attachment2)
    environment = EnvironmentSeries(
        environment_frame.iloc[:, 0].to_numpy(dtype=float),
        environment_frame.iloc[:, 1].to_numpy(dtype=float),
        environment_frame.iloc[:, 2].to_numpy(dtype=float),
    )
    radius = RadiusSeries(
        radius_frame.iloc[:, 0].to_numpy(dtype=float),
        radius_frame.iloc[:, 1].to_numpy(dtype=float),
        "linear",
    )
    return environment, radius


def timed_bdf(
    label: str,
    environment: EnvironmentSeries,
    radius: RadiusSeries,
    config: SolverConfig,
    post_boundary: PostBoundary | None = None,
) -> tuple[SimulationResult, float]:
    print(f"START BDF {label}", flush=True)
    started = time.perf_counter()
    result = solve_problem4_bdf(environment, radius, config, post_boundary)
    runtime = time.perf_counter() - started
    print(f"DONE  BDF {label}: {result.drying_event_time_s / 3600.0:.9f} h ({runtime:.2f} s)", flush=True)
    return result, runtime


def timed_be(
    label: str,
    environment: EnvironmentSeries,
    radius: RadiusSeries,
    config: BESolverConfig,
    post_boundary: PostBoundary | None = None,
) -> tuple[BEResult, float]:
    print(f"START BE  {label}", flush=True)
    started = time.perf_counter()
    result = solve_problem4_be(environment, radius, config, post_boundary)
    runtime = time.perf_counter() - started
    print(f"DONE  BE  {label}: {result.drying_event_time_s / 3600.0:.9f} h ({runtime:.2f} s)", flush=True)
    return result, runtime


def bdf_summary(result: SimulationResult, runtime_s: float) -> dict:
    return {
        "node_count": result.diagnostics.internal_node_count,
        "event_time_s": result.drying_event_time_s,
        "event_time_h": result.drying_event_time_s / 3600.0,
        "report_end_time_s": result.report_end_time_s,
        "event_radius_cm": result.event_radius_cm,
        "event_center_moisture": float(result.event_moisture_internal[0]),
        "event_surface_moisture": float(result.event_moisture_internal[-1]),
        "event_argmax_xi": float(np.argmax(result.event_moisture_internal) / (len(result.event_moisture_internal) - 1)),
        "runtime_s": runtime_s,
        "diagnostics": asdict(result.diagnostics),
    }


def be_summary(result: BEResult, runtime_s: float) -> dict:
    return {
        "node_count": len(result.event_moisture_internal),
        "event_time_s": result.drying_event_time_s,
        "event_time_h": result.drying_event_time_s / 3600.0,
        "event_radius_cm": result.event_radius_cm,
        "event_center_moisture": float(result.event_moisture_internal[0]),
        "event_surface_moisture": float(result.event_moisture_internal[-1]),
        "event_argmax_xi": float(np.argmax(result.event_moisture_internal) / (len(result.event_moisture_internal) - 1)),
        "runtime_s": runtime_s,
        "diagnostics": asdict(result.diagnostics),
    }


def common_profile_difference(
    left_times: np.ndarray,
    left_values: np.ndarray,
    right_times: np.ndarray,
    right_values: np.ndarray,
) -> float:
    common = sorted(set(left_times.astype(int)).intersection(right_times.astype(int)))
    common = [value for value in common if value % 21600 == 0]
    if not common:
        return float("inf")
    maximum = 0.0
    for time_s in common:
        left = left_values[int(np.where(np.isclose(left_times, time_s))[0][0])]
        right = right_values[int(np.where(np.isclose(right_times, time_s))[0][0])]
        finite = np.isfinite(left) & np.isfinite(right)
        if np.any(finite):
            maximum = max(maximum, float(np.max(np.abs(left[finite] - right[finite]))))
    return maximum


def compare_bdf(left: SimulationResult, right: SimulationResult) -> dict:
    difference_s = abs(left.drying_event_time_s - right.drying_event_time_s)
    return {
        "event_time_difference_s": difference_s,
        "event_time_relative_difference": difference_s / right.drying_event_time_s,
        "paper_sample_maximum_moisture_difference": common_profile_difference(
            left.time_s, left.moisture, right.time_s, right.moisture
        ),
        "event_location_same": bool(
            np.argmax(left.event_moisture_internal) == 0 and np.argmax(right.event_moisture_internal) == 0
        ),
    }


def compare_bdf_be(left: SimulationResult, right: BEResult) -> dict:
    difference_s = abs(left.drying_event_time_s - right.drying_event_time_s)
    return {
        "event_time_difference_s": difference_s,
        "event_time_relative_difference": difference_s / left.drying_event_time_s,
        "paper_sample_maximum_moisture_difference": common_profile_difference(
            left.time_s, left.moisture, right.time_s, right.moisture
        ),
    }


def write_result_xlsx(result: SimulationResult, template: Path, out_path: Path) -> None:
    """按附件3模板格式写出全截面水分浓度剖面。

    全精度存储、仅用 number_format 控制四位小数显示（与原 result4.xlsx 一致）；
    半径收缩后已退出药材内部的固定距离留空（NaN -> 空单元格）。
    """
    workbook = openpyxl.load_workbook(template)
    sheet = workbook.worksheets[0]
    fixed = result.fixed_distance_cm
    sheet.cell(1, 1).value = "时间\\到药材中心的距离"
    for j in range(fixed.size):
        sheet.cell(1, 2 + j).value = round(float(fixed[j]), 10)
    sheet.cell(1, 2 + fixed.size).value = "药材表面"
    for i in range(result.time_s.size):
        time_cell = sheet.cell(i + 2, 1)
        time_cell.value = int(round(result.time_s[i]))
        time_cell.number_format = "0"
        values = result.moisture[i]
        for j in range(values.size):
            cell = sheet.cell(i + 2, 2 + j)
            value = float(values[j])
            cell.value = value if np.isfinite(value) else None
            cell.number_format = "0.0000"
    sheet.column_dimensions["A"].width = 24
    for j in range(fixed.size):
        sheet.column_dimensions[get_column_letter(2 + j)].width = 10
    sheet.column_dimensions[get_column_letter(2 + fixed.size)].width = 14
    sheet.freeze_panes = "A2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def main() -> None:
    missing = [str(path) for path in (ATTACHMENT1, ATTACHMENT2, TEMPLATE) if not path.exists()]
    if missing:
        raise SystemExit(
            "缺少附件，无法运行问题四：\n  "
            + "\n  ".join(missing)
            + "\n请将 附件/（附件1.xlsx、附件2.xlsx、附件3/result4.xlsx）放在与本文件同一根目录。"
        )
    environment, radius = load_inputs(ATTACHMENT1, ATTACHMENT2)

    base_bdf = SolverConfig(
        node_count=201,
        sample_interval_s=21600,
        relative_tolerance=1.0e-8,
        temperature_absolute_tolerance=1.0e-8,
        moisture_absolute_tolerance=1.0e-10,
        max_time_step_s=30.0,
    )

    # 空间网格逐级加密（物质坐标节点数 N），观察相邻差收敛
    spatial_node_counts = [201, 401, 801, 1601]
    spatial_results: dict[int, SimulationResult] = {}
    spatial_rows = []
    previous_event = None
    for node_count in spatial_node_counts:
        result, runtime = timed_bdf(
            f"spatial n={node_count}", environment, radius, replace(base_bdf, node_count=node_count)
        )
        spatial_results[node_count] = result
        row = bdf_summary(result, runtime)
        row["difference_from_previous_s"] = (
            None if previous_event is None else result.drying_event_time_s - previous_event
        )
        spatial_rows.append(row)
        previous_event = result.drying_event_time_s

    def _accepted(metrics: dict) -> bool:
        return bool(
            metrics["event_time_difference_s"] < 60.0
            and metrics["event_time_relative_difference"] < 1.0e-3
            and metrics["paper_sample_maximum_moisture_difference"] < 5.0e-5
            and metrics["event_location_same"]
        )

    spatial_pairs = []
    for coarse, fine in zip(spatial_node_counts[:-1], spatial_node_counts[1:], strict=True):
        metrics = compare_bdf(spatial_results[coarse], spatial_results[fine])
        metrics.update({"coarse_node_count": coarse, "fine_node_count": fine})
        metrics["accepted"] = _accepted(metrics)
        spatial_pairs.append(metrics)
    if not spatial_pairs[-1]["accepted"]:
        node_count = 3201
        result, runtime = timed_bdf(
            f"spatial n={node_count}", environment, radius, replace(base_bdf, node_count=node_count)
        )
        spatial_results[node_count] = result
        row = bdf_summary(result, runtime)
        row["difference_from_previous_s"] = result.drying_event_time_s - previous_event
        spatial_rows.append(row)
        metrics = compare_bdf(spatial_results[1601], result)
        metrics.update({"coarse_node_count": 1601, "fine_node_count": 3201})
        metrics["accepted"] = _accepted(metrics)
        spatial_pairs.append(metrics)
        spatial_node_counts.append(node_count)
    selected_node_count = spatial_node_counts[-1]
    if not spatial_pairs[-1]["accepted"]:
        raise RuntimeError(f"finest spatial pair did not pass: {spatial_pairs[-1]}")

    # BDF 时间容差 / 步长三档
    bdf_configs = {
        "B1": replace(
            base_bdf,
            node_count=selected_node_count,
            relative_tolerance=1.0e-7,
            temperature_absolute_tolerance=1.0e-7,
            moisture_absolute_tolerance=1.0e-9,
            max_time_step_s=60.0,
        ),
        "B2": replace(base_bdf, node_count=selected_node_count),
        "B3": replace(
            base_bdf,
            node_count=selected_node_count,
            relative_tolerance=1.0e-9,
            temperature_absolute_tolerance=1.0e-10,
            moisture_absolute_tolerance=1.0e-12,
            max_time_step_s=10.0,
        ),
    }
    bdf_results: dict[str, SimulationResult] = {"B2": spatial_results[selected_node_count]}
    bdf_runtime = {"B2": spatial_rows[-1]["runtime_s"]}
    for name in ("B1", "B3"):
        bdf_results[name], bdf_runtime[name] = timed_bdf(
            f"time {name}", environment, radius, bdf_configs[name]
        )
    bdf_rows = [
        {"name": name, "config": asdict(bdf_configs[name]), **bdf_summary(bdf_results[name], bdf_runtime[name])}
        for name in ("B1", "B2", "B3")
    ]
    bdf_comparisons = {
        "B1_vs_B2": compare_bdf(bdf_results["B1"], bdf_results["B2"]),
        "B2_vs_B3": compare_bdf(bdf_results["B2"], bdf_results["B3"]),
    }
    bdf_pass = bool(
        bdf_comparisons["B2_vs_B3"]["event_time_difference_s"] < 10.0
        and bdf_comparisons["B2_vs_B3"]["paper_sample_maximum_moisture_difference"] < 1.0e-4
    )

    # 后向欧拉-Picard 对照：dt=20/10/5/2.5 s 逐档收敛
    be_results: dict[float, BEResult] = {}
    be_rows = []
    for dt_s in (20.0, 10.0, 5.0, 2.5):
        config = BESolverConfig(node_count=selected_node_count, time_step_s=dt_s, sample_interval_s=21600)
        result, runtime = timed_be(f"time dt={dt_s:g}s", environment, radius, config)
        be_results[dt_s] = result
        be_rows.append({"dt_s": dt_s, "config": asdict(config), **be_summary(result, runtime)})
    be_pairs = [
        {
            "coarse_dt_s": coarse,
            "fine_dt_s": fine,
            "event_time_difference_s": abs(be_results[coarse].drying_event_time_s - be_results[fine].drying_event_time_s),
            "paper_sample_maximum_moisture_difference": common_profile_difference(
                be_results[coarse].time_s,
                be_results[coarse].moisture,
                be_results[fine].time_s,
                be_results[fine].moisture,
            ),
        }
        for coarse, fine in ((20.0, 10.0), (10.0, 5.0), (5.0, 2.5))
    ]
    be_pass = bool(
        be_pairs[-1]["event_time_difference_s"] < 30.0
        and be_pairs[-1]["paper_sample_maximum_moisture_difference"] < 1.0e-4
        and all(row["diagnostics"]["unconverged_steps"] == 0 for row in be_rows)
    )
    cross_solver = compare_bdf_be(bdf_results["B3"], be_results[2.5])
    cross_pass = bool(
        cross_solver["event_time_difference_s"] < 60.0
        and cross_solver["paper_sample_maximum_moisture_difference"] < 1.0e-4
    )

    # 后期边界情景敏感性（缓存闭包：默认末值 + hm×1.0 复用主运行）
    main_reference = bdf_results["B2"]
    default_post = post_boundary_from_tail(environment)
    scenario_config = bdf_configs["B2"]
    cache: dict[tuple, tuple[SimulationResult, float]] = {
        (default_post.temperature_c, default_post.moisture, 1.0): (main_reference, bdf_runtime["B2"])
    }

    def scenario_run(name: str, post: PostBoundary, hm_factor: float) -> dict:
        key = (round(post.temperature_c, 12), round(post.moisture, 12), hm_factor)
        if key not in cache:
            cache[key] = timed_bdf(
                f"boundary {name}",
                environment,
                radius,
                replace(scenario_config, mass_transfer_m_s=scenario_config.mass_transfer_m_s * hm_factor),
                post,
            )
        result, runtime = cache[key]
        return {
            "scenario": name,
            "post_temperature_c": post.temperature_c,
            "post_moisture": post.moisture,
            "mass_transfer_factor": hm_factor,
            **bdf_summary(result, runtime),
            "difference_from_default_s": result.drying_event_time_s - main_reference.drying_event_time_s,
        }

    data_rows = [
        scenario_run("attachment_terminal_value", default_post, 1.0),
        scenario_run("last_30_min_time_average", post_boundary_from_tail(environment, 1800.0), 1.0),
        scenario_run("last_1_h_time_average", post_boundary_from_tail(environment, 3600.0), 1.0),
    ]
    engineering_rows = [
        scenario_run(item.name, item.post_boundary, item.mass_transfer_factor)
        for item in engineering_boundary_scenarios(default_post)
    ]
    hm_rows = [
        scenario_run(item.name, item.post_boundary, item.mass_transfer_factor)
        for item in mass_transfer_scenarios(default_post)
    ]
    boundary_times = [row["event_time_s"] for row in data_rows + engineering_rows + hm_rows]

    # 收缩消融：固定半径 2 cm 反事实对照
    fixed_result, fixed_runtime = timed_bdf(
        "fixed radius 2 cm",
        environment,
        radius,
        replace(
            scenario_config,
            fixed_radius_cm=2.0,
            initial_time_limit_s=120.0 * 3600.0,
            fallback_time_limit_s=168.0 * 3600.0,
        ),
    )
    # 坐标模型形式敏感性：Euler 空间坐标（保留收缩对流项）
    eulerian_result, eulerian_runtime = timed_bdf(
        "Eulerian model form", environment, radius, replace(scenario_config, model_form="eulerian")
    )

    # 最终 60 s 输出运行（连续解与主运行一致，仅采样间隔不同）
    print("START final 60 s output run", flush=True)
    final_result, final_runtime = timed_bdf(
        "final output", environment, radius, replace(scenario_config, sample_interval_s=60)
    )
    if abs(final_result.drying_event_time_s - main_reference.drying_event_time_s) > 1.0e-6:
        raise RuntimeError("sampling interval changed the continuous solution")

    quality_pass = bool(
        final_result.diagnostics.maximum_normalized_balance_residual < 1.0e-6
        and final_result.diagnostics.minimum_raw_moisture >= -1.0e-8
        and final_result.diagnostics.radial_monotonicity_all_outputs
        and np.argmax(final_result.event_moisture_internal) == 0
        and final_result.previous_report_max_moisture >= 0.15 - 1.0e-10
        and final_result.report_max_moisture < 0.15
    )

    verification = {
        "input": {
            "environment_terminal_time_s": environment.terminal_time_s,
            "environment_terminal_temperature_c": environment.terminal_temperature_c,
            "environment_terminal_moisture": environment.terminal_moisture,
            "radius_terminal_time_s": radius.terminal_time_s,
            "radius_initial_cm": float(radius.radius_cm[0]),
            "radius_terminal_cm": float(radius.radius_cm[-1]),
        },
        "spatial_grid": {
            "runs": spatial_rows,
            "adjacent_comparisons": spatial_pairs,
            "selected_node_count": selected_node_count,
            "passed": bool(spatial_pairs[-1]["accepted"]),
            "oscillatory": bool(
                len(spatial_rows) >= 3
                and np.sign(spatial_rows[-1]["difference_from_previous_s"])
                != np.sign(spatial_rows[-2]["difference_from_previous_s"])
            ),
        },
        "bdf_time_sensitivity": {"runs": bdf_rows, "comparisons": bdf_comparisons, "passed": bdf_pass},
        "be_time_sensitivity": {"runs": be_rows, "comparisons": be_pairs, "passed": be_pass},
        "cross_solver": {"comparison": cross_solver, "passed": cross_pass},
        "boundary_sensitivity": {
            "data_driven": data_rows,
            "engineering_3x3": engineering_rows,
            "mass_transfer": hm_rows,
            "all_scenarios_range_s": [float(min(boundary_times)), float(max(boundary_times))],
            "all_scenarios_range_h": [float(min(boundary_times) / 3600.0), float(max(boundary_times) / 3600.0)],
        },
        "shrinkage_ablation": {
            "fixed_radius_event_time_s": fixed_result.drying_event_time_s,
            "fixed_radius_event_time_h": fixed_result.drying_event_time_s / 3600.0,
            "dynamic_radius_event_time_s": main_reference.drying_event_time_s,
            "dynamic_radius_event_time_h": main_reference.drying_event_time_s / 3600.0,
            "time_reduction_s": fixed_result.drying_event_time_s - main_reference.drying_event_time_s,
            "time_reduction_h": (fixed_result.drying_event_time_s - main_reference.drying_event_time_s) / 3600.0,
            "fixed_radius_runtime_s": fixed_runtime,
        },
        "model_form_sensitivity": {
            "material_event_time_s": main_reference.drying_event_time_s,
            "material_event_time_h": main_reference.drying_event_time_s / 3600.0,
            "eulerian_event_time_s": eulerian_result.drying_event_time_s,
            "eulerian_event_time_h": eulerian_result.drying_event_time_s / 3600.0,
            "absolute_difference_s": abs(eulerian_result.drying_event_time_s - main_reference.drying_event_time_s),
            "absolute_difference_h": abs(eulerian_result.drying_event_time_s - main_reference.drying_event_time_s) / 3600.0,
            "eulerian_runtime_s": eulerian_runtime,
        },
        "final": {
            "all_numerical_acceptance_checks_passed": bool(
                spatial_pairs[-1]["accepted"] and bdf_pass and be_pass and cross_pass and quality_pass
            ),
            "selected_node_count": selected_node_count,
            "bdf_configuration": "B2",
            "continuous_event_time_s": final_result.drying_event_time_s,
            "continuous_event_time_h": final_result.drying_event_time_s / 3600.0,
            "first_strict_60s_time_s": final_result.report_end_time_s,
            "first_strict_60s_time_h": final_result.report_end_time_s / 3600.0,
            "previous_report_time_s": final_result.previous_report_time_s,
            "previous_max_moisture": final_result.previous_report_max_moisture,
            "event_max_moisture": float(np.max(final_result.event_moisture_internal)),
            "first_strict_report_max_moisture": final_result.report_max_moisture,
            "event_center_moisture": float(final_result.event_moisture_internal[0]),
            "event_surface_moisture": float(final_result.event_moisture_internal[-1]),
            "event_radius_cm": final_result.event_radius_cm,
            "argmax_location_xi": float(np.argmax(final_result.event_moisture_internal) / (selected_node_count - 1)),
            "minimum_raw_moisture": final_result.diagnostics.minimum_raw_moisture,
            "clip_diagnostics": {
                "clipped_cell_count": final_result.diagnostics.clipped_cell_count,
                "maximum_clip_magnitude": final_result.diagnostics.maximum_clip_magnitude,
            },
            "maximum_normalized_balance_residual": final_result.diagnostics.maximum_normalized_balance_residual,
            "dry_mass_index_relative_range": final_result.diagnostics.dry_mass_index_relative_range,
            "boundary_all_scenarios_range_h": [
                float(min(boundary_times) / 3600.0),
                float(max(boundary_times) / 3600.0),
            ],
            "final_output_runtime_s": final_runtime,
            "quality_checks_passed": quality_pass,
        },
    }
    json_dump(OUT_JSON, verification)
    write_result_xlsx(final_result, TEMPLATE, OUT_XLSX)

    def _ok(passed: bool) -> str:
        return "通过" if passed else "未通过"

    final = verification["final"]
    boundary = verification["boundary_sensitivity"]
    bar = "=" * 64
    print(bar, flush=True)
    print("问题四 严格物质坐标 FVM+BDF —— 运行与验证结果汇总", flush=True)
    print(bar, flush=True)
    print(f"[主答案] 连续临界时间 t* = {final['continuous_event_time_h']:.6f} h ({final['continuous_event_time_s']:.3f} s)", flush=True)
    print(f"[主答案] 首个严格达标 60 s 时刻 = {final['first_strict_60s_time_h']:.6f} h", flush=True)
    print(f"[主答案] 后期边界全部情景范围 = {boundary['all_scenarios_range_h'][0]:.4f}-{boundary['all_scenarios_range_h'][1]:.4f} h", flush=True)
    print("-" * 64, flush=True)
    print("[空间网格收敛]", flush=True)
    for row in verification["spatial_grid"]["runs"]:
        diff = row.get("difference_from_previous_s")
        diff_s = "--" if diff is None else f"{diff:+.3f} s"
        print(f"  N={row['node_count']:>5}: t*={row['event_time_h']:.9f} h  (相邻差 {diff_s}, {row['runtime_s']:.1f} s)", flush=True)
    print(f"  相邻网格验收：{_ok(verification['spatial_grid']['passed'])}", flush=True)
    print("[BDF 容差/步长敏感性]", flush=True)
    for row in verification["bdf_time_sensitivity"]["runs"]:
        print(f"  {row['name']}: rtol={row['config']['relative_tolerance']:.0e} max_step={row['config']['max_time_step_s']:g}s -> t*={row['event_time_h']:.9f} h", flush=True)
    print(f"  BDF 时间敏感性：{_ok(verification['bdf_time_sensitivity']['passed'])}", flush=True)
    print("[后向欧拉交叉验证]", flush=True)
    for row in verification["be_time_sensitivity"]["runs"]:
        print(f"  dt={row['dt_s']:g}s -> t*={row['event_time_h']:.9f} h (Picard max {row['diagnostics']['maximum_picard_iterations_used']}, 未收敛步 {row['diagnostics']['unconverged_steps']})", flush=True)
    print(f"  BE 时间敏感性：{_ok(verification['be_time_sensitivity']['passed'])}；BDF/BE 交叉：{_ok(verification['cross_solver']['passed'])}（差 {cross_solver['event_time_difference_s']:.3f} s）", flush=True)
    print("[后期边界情景敏感性]", flush=True)
    for row in boundary["data_driven"] + boundary["engineering_3x3"] + boundary["mass_transfer"]:
        print(f"  {row['scenario']}: T={row['post_temperature_c']:.4f}C C={row['post_moisture']:.4f} hm x{row['mass_transfer_factor']:.1f} -> t*={row['event_time_h']:.6f} h ({row['difference_from_default_s']/60.0:+.2f} min)", flush=True)
    sh = verification["shrinkage_ablation"]
    print(f"[收缩消融] 动态半径={sh['dynamic_radius_event_time_h']:.6f} h  固定2cm={sh['fixed_radius_event_time_h']:.6f} h  缩短={sh['time_reduction_h']:.6f} h", flush=True)
    mf = verification["model_form_sensitivity"]
    print(f"[坐标模型形式] 物质坐标={mf['material_event_time_h']:.6f} h  Euler={mf['eulerian_event_time_h']:.6f} h  差={mf['absolute_difference_h']:.6f} h", flush=True)
    print("-" * 64, flush=True)
    print(f"[守恒核验] 最大归一化通量残差={final['maximum_normalized_balance_residual']:.3e}  最小原始含水率={final['minimum_raw_moisture']:.9f} kg/kg", flush=True)
    print(f"[干物质诊断] 相对变幅={final['dry_mass_index_relative_range']:.2%}（模型局限指标，非失败判据）", flush=True)
    print(f"[停机判据] 上一报告时刻最大含水率={final['previous_max_moisture']:.9f}（>=0.15）  首个严格达标时刻={final['first_strict_report_max_moisture']:.9f}（<0.15）  argmax ξ={final['argmax_location_xi']:.4f}", flush=True)
    print(f"[总验收] 全部数值验收：{_ok(final['all_numerical_acceptance_checks_passed'])}", flush=True)
    print(bar, flush=True)
    print(f"已写出：{OUT_XLSX.name}（全截面水分浓度剖面）、{OUT_JSON.name}（结构化验证，兼画图数据源）", flush=True)


if __name__ == "__main__":
    main()
