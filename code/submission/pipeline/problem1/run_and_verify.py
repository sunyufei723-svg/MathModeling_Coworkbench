"""problem1 改进版：运行 + 数值复核 + 结果生成（独立目录，不动 code/problem1 与 results/problem1 原文件）。

一条命令完成：
  1. 读附件1（openpyxl，绕过依赖 @oai/artifact-tool 的 extract_environment.mjs）；
  2. 跑三种配置：baseline（算术平均、无裁剪＝原版行为）/ improved（谐波平均＋C≥0 裁剪）/
     grid-refined improved（内部网格加密一倍，做网格无关性）；
  3. 数值复核：baseline 复现原 result1.xlsx（验证忠实拷贝）、改进效果（harmonic vs arithmetic，
     同环境纯归因）、温度场解耦一致性、网格无关性、C≥0 物理核验；
  4. 生成 results/problem1_improved/：result1.xlsx（copy 原版保格式 + 写改进数据）+
     problem1_improved_tables.md + verification.json，并打印精简摘要。

numerical_verification.md 依据 verification.json 手写（与 agent_C 的 problem2_fvm_bdf 一致）。
"""

from __future__ import annotations

import argparse
import json
import os
import shutil
import sys
from dataclasses import asdict, replace
from pathlib import Path

import numpy as np
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "business" / "problem1"))

from solver import EnvironmentSeries, SimulationResult, SolverConfig, solve_problem1


PROJECT_ROOT = Path(__file__).resolve().parents[4]
ORIGINAL_RESULT = PROJECT_ROOT / "results" / "problem1" / "result1.xlsx"
DEFAULT_OUT_DIR = PROJECT_ROOT / "results" / "problem1_improved"
SAMPLE_TIMES = np.array([100, 300, 600, 900, 1200, 1500, 1800], dtype=float)
SAMPLE_RADII_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0], dtype=float)
TEMP_SHEET = "温度"
MOIST_SHEET = "水分浓度"


def resolve_attachment1(explicit_path: Path | None = None) -> Path:
    """定位附件1.xlsx（本地/网盘流转，未入库），可用 --attachment1 或 CUMCM_A_PROBLEM_DIR 指定。"""

    if explicit_path is not None:
        candidates = [explicit_path]
    else:
        candidates = []
        source_dir = os.environ.get("CUMCM_A_PROBLEM_DIR")
        if source_dir:
            candidates.append(Path(source_dir) / "附件" / "附件1.xlsx")
        candidates.append(PROJECT_ROOT / "附件" / "附件1.xlsx")
    for candidate in candidates:
        if candidate.exists():
            return candidate
    searched = "\n".join(str(path) for path in candidates)
    raise FileNotFoundError(f"未找到附件1.xlsx，请用 --attachment1 指定或设置 CUMCM_A_PROBLEM_DIR。\n已搜索：\n{searched}")


def load_environment(path: Path) -> EnvironmentSeries:
    """openpyxl 读附件1 三列（时间/温度/水分浓度）→ EnvironmentSeries。"""

    worksheet = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    rows = [r for r in worksheet.iter_rows(min_row=2, values_only=True) if r[0] not in (None, "")]
    return EnvironmentSeries(
        time_s=np.asarray([float(r[0]) for r in rows], dtype=float),
        temperature_c=np.asarray([float(r[1]) for r in rows], dtype=float),
        moisture=np.asarray([float(r[2]) for r in rows], dtype=float),
    )


def selected_values(result: SimulationResult, values: np.ndarray) -> np.ndarray:
    time_indices = [int(np.where(np.isclose(result.time_s, value))[0][0]) for value in SAMPLE_TIMES]
    radius_indices = [int(np.where(np.isclose(result.radius_cm, value))[0][0]) for value in SAMPLE_RADII_CM]
    return values[np.ix_(time_indices, radius_indices)]


def difference_summary(left: np.ndarray, right: np.ndarray, times: np.ndarray, radii: np.ndarray) -> dict[str, object]:
    """两数组逐点对照：最大绝对差 + 位置 + 平均 + 四位小数一致格数。"""

    difference = left - right
    absolute = np.abs(difference)
    location = np.unravel_index(int(np.argmax(absolute)), absolute.shape)
    return {
        "maximum_absolute_difference": float(absolute[location]),
        "signed_difference_at_maximum": float(difference[location]),
        "time_s_at_maximum": float(times[location[0]]),
        "radius_cm_at_maximum": float(radii[location[1]]),
        "mean_absolute_difference": float(absolute.mean()),
        "four_decimal_cells_equal": int(np.count_nonzero(np.round(left, 4) == np.round(right, 4))),
        "cell_count": int(left.size),
    }


def load_original_result(path: Path) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    """openpyxl 读原 result1.xlsx 的温度/水分两表（行=t1..1800，列=r0..2.0），返回 (time, temp, moist)。"""

    workbook = openpyxl.load_workbook(path, data_only=True)
    sheet_t = workbook[TEMP_SHEET]
    sheet_m = workbook[MOIST_SHEET]
    n_rows = sheet_t.max_row
    n_cols = sheet_t.max_column
    time_s = np.asarray([sheet_t.cell(row=r, column=1).value for r in range(2, n_rows + 1)], dtype=float)
    temperature = np.asarray(
        [[sheet_t.cell(row=r, column=c).value for c in range(2, n_cols + 1)] for r in range(2, n_rows + 1)],
        dtype=float,
    )
    moisture = np.asarray(
        [[sheet_m.cell(row=r, column=c).value for c in range(2, n_cols + 1)] for r in range(2, n_rows + 1)],
        dtype=float,
    )
    return time_s, temperature, moisture


def write_result_xlsx(result: SimulationResult, out_path: Path) -> None:
    """copy 原版 result1.xlsx（保留表头/列宽/number_format）→ 只改写数据区（行=t1..1800，列=r0..2.0）。"""

    out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(ORIGINAL_RESULT, out_path)
    workbook = openpyxl.load_workbook(out_path)
    sheet_t = workbook[TEMP_SHEET]
    sheet_m = workbook[MOIST_SHEET]
    n_radius = result.radius_cm.size
    for time_index in range(1, result.time_s.size):  # t=1..1800（原版模板从 t=1 起）
        row = time_index + 1
        for radius_index in range(n_radius):
            column = radius_index + 2
            sheet_t.cell(row=row, column=column).value = round(float(result.temperature_c[time_index, radius_index]), 10)
            sheet_m.cell(row=row, column=column).value = round(float(result.moisture[time_index, radius_index]), 10)
    workbook.save(out_path)


def markdown_table(headers: list[str], rows: list[list[float]]) -> str:
    lines = [
        "| " + " | ".join(headers) + " |",
        "| " + " | ".join(["---:"] * len(headers)) + " |",
    ]
    for row in rows:
        cells = [str(int(row[0]))] + [f"{value:.4f}" for value in row[1:]]
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def _table_rows(result: SimulationResult, values: np.ndarray) -> list[list[float]]:
    rows: list[list[float]] = []
    for sample_time in SAMPLE_TIMES:
        time_index = int(np.where(np.isclose(result.time_s, sample_time))[0][0])
        row = [float(sample_time)]
        for radius in SAMPLE_RADII_CM:
            radius_index = int(np.where(np.isclose(result.radius_cm, radius))[0][0])
            row.append(float(values[time_index, radius_index]))
        rows.append(row)
    return rows


def write_tables(result: SimulationResult, out_path: Path) -> None:
    headers = ["时间/s", "0 cm", "0.5 cm", "1 cm", "1.5 cm", "2 cm"]
    text = [
        "# A题第一问改进版论文表格结果",
        "",
        "> 计算方法：一维径向守恒型有限体积法 + 隐式 BDF；内部网格 0.0025 cm；界面扩散系数谐波平均；",
        "> 输出前水分 C≥0 裁剪。与原版 `results/problem1/` 并存，未覆盖原结果。",
        "",
        "## 表1 30分钟内药材的温度 / °C",
        "",
        markdown_table(headers, _table_rows(result, result.temperature_c)),
        "",
        "## 表2 30分钟内药材的水分浓度 / kg·kg⁻¹",
        "",
        markdown_table(headers, _table_rows(result, result.moisture)),
        "",
    ]
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text("\n".join(text), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="运行 problem1 改进解并完成数值复核")
    parser.add_argument("--attachment1", type=Path, default=None, help="附件1.xlsx 路径（默认自动定位）")
    parser.add_argument("--out-dir", type=Path, default=DEFAULT_OUT_DIR, help="结果输出目录")
    parser.add_argument("--skip-grid-refine", action="store_true", help="跳过网格加密复核（加快运行）")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    out_dir: Path = args.out_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    attachment = resolve_attachment1(args.attachment1)
    environment = load_environment(attachment)

    baseline_config = SolverConfig(interface_mean="arithmetic", clip_negative_moisture=False)  # ＝原版行为
    improved_config = SolverConfig()  # 谐波平均 + C≥0 裁剪
    grid_config = replace(improved_config, internal_dr_cm=improved_config.internal_dr_cm / 2.0)

    print("running baseline (arithmetic, no clip)", flush=True)
    baseline = solve_problem1(environment, baseline_config)
    print("running improved (harmonic, clip)", flush=True)
    improved = solve_problem1(environment, improved_config)
    grid = None
    if not args.skip_grid_refine:
        print(f"running grid-refined improved (dr={grid_config.internal_dr_cm} cm)", flush=True)
        grid = solve_problem1(environment, grid_config)

    original_time, original_temp, original_moist = load_original_result(ORIGINAL_RESULT)
    if not np.array_equal(original_time, improved.time_s[1:]):
        raise ValueError("原 result1.xlsx 的时间列与改进解导出时间网格不一致")

    verification = {
        "source_attachment": str(attachment),
        "configs": {
            "baseline": asdict(baseline_config),
            "improved": asdict(improved_config),
            "grid_refined": asdict(grid_config) if grid is not None else None,
        },
        "baseline_reproduction_vs_original": {
            "note": "arithmetic+no-clip（原版行为）对 C 的 result1.xlsx；差异源于 scipy 积分器版本，非流程错误",
            "temperature_c": difference_summary(baseline.temperature_c[1:], original_temp, baseline.time_s[1:], baseline.radius_cm),
            "moisture": difference_summary(baseline.moisture[1:], original_moist, baseline.time_s[1:], baseline.radius_cm),
        },
        "improvement_effect_harmonic_vs_arithmetic": {
            "note": "improved vs baseline，同一 scipy 环境，唯一变量＝界面平均方式＋裁剪 → 纯改进归因",
            "temperature_c": difference_summary(improved.temperature_c, baseline.temperature_c, improved.time_s, improved.radius_cm),
            "moisture": difference_summary(improved.moisture, baseline.moisture, improved.time_s, improved.radius_cm),
        },
        "improved_vs_original": {
            "note": "改进解对 C 原 result1.xlsx 的总差异（＝改进效果 + scipy 版本噪声）",
            "temperature_c": difference_summary(improved.temperature_c[1:], original_temp, improved.time_s[1:], improved.radius_cm),
            "moisture": difference_summary(improved.moisture[1:], original_moist, improved.time_s[1:], improved.radius_cm),
        },
        "physical_checks": {
            "improved_moisture_min": float(improved.moisture.min()),
            "baseline_moisture_min": float(baseline.moisture.min()),
            "clip_is_noop": bool(baseline.moisture.min() > 0.0),
            "temperature_identical_baseline_vs_improved": bool(np.array_equal(improved.temperature_c, baseline.temperature_c)),
            "improved_all_finite": bool(np.isfinite(improved.temperature_c).all() and np.isfinite(improved.moisture).all()),
        },
        "selected_tables": {
            "temperature_c": [[float(t), *[float(v) for v in row]] for t, row in zip(SAMPLE_TIMES, selected_values(improved, improved.temperature_c))],
            "moisture": [[float(t), *[float(v) for v in row]] for t, row in zip(SAMPLE_TIMES, selected_values(improved, improved.moisture))],
        },
    }
    if grid is not None:
        verification["grid_independence"] = {
            "note": f"改进解 dr={improved_config.internal_dr_cm} cm vs {grid_config.internal_dr_cm} cm（同一输出网格）",
            "temperature_c": difference_summary(improved.temperature_c, grid.temperature_c, improved.time_s, improved.radius_cm),
            "moisture": difference_summary(improved.moisture, grid.moisture, improved.time_s, improved.radius_cm),
        }

    write_result_xlsx(improved, out_dir / "result1.xlsx")
    write_tables(improved, out_dir / "problem1_improved_tables.md")
    (out_dir / "verification.json").write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")

    compact = {
        "out_dir": str(out_dir),
        "baseline_reproduction": verification["baseline_reproduction_vs_original"],
        "improvement_effect": verification["improvement_effect_harmonic_vs_arithmetic"],
        "physical_checks": verification["physical_checks"],
        "grid_independence": verification.get("grid_independence"),
    }
    print(json.dumps(compact, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
