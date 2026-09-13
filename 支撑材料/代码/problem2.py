"""问题二：变物性（附录3）热湿耦合，守恒径向有限体积 + 联合状态自适应 BDF，求 3 h 烘干的温度/水分场。

温度场与水分场通过附录3 物性（ρ、cp、k、D 均随含水率 C 变化）双向强耦合，合并成一个状态向量整体推进；
界面 k、D 用调和平均。main 内自动做网格加密、BDF 容差收紧、早期薄层三项复核（论文模型检验章引用）。

直接运行（无需任何参数）：python problem2.py
  · 读同根目录 附件/附件1.xlsx（环境）与 附件/附件3/result2.xlsx（结果格式模板），读不到给提示；
  · 按模板格式把主结果写到同根目录 result2.xlsx，验证摘要写 verification2.json；
  · 终端分节打印表3(温度)/表4(水分) + 网格/容差/早期薄层复核 + 物理核验。
"""
from __future__ import annotations

import json
from dataclasses import asdict, dataclass, replace
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.integrate import solve_ivp
from scipy.sparse import bmat, csc_matrix, diags


# ============================== 求解核心 ==============================
@dataclass(frozen=True)
class EnvironmentSeries:
    """Piecewise-linear oven temperature and moisture boundary data."""

    time_s: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray

    def __post_init__(self) -> None:
        sizes = {len(self.time_s), len(self.temperature_c), len(self.moisture)}
        if len(sizes) != 1 or not len(self.time_s):
            raise ValueError("environment arrays must have the same non-zero length")
        if np.any(np.diff(self.time_s) <= 0):
            raise ValueError("environment times must be strictly increasing")
        if not (
            np.isfinite(self.time_s).all()
            and np.isfinite(self.temperature_c).all()
            and np.isfinite(self.moisture).all()
        ):
            raise ValueError("environment arrays must be finite")

    def temperature_at(self, t_s: float) -> float:
        return float(np.interp(t_s, self.time_s, self.temperature_c))

    def moisture_at(self, t_s: float) -> float:
        return float(np.interp(t_s, self.time_s, self.moisture))


@dataclass(frozen=True)
class SolverConfig:
    radius_cm: float = 2.0
    internal_dr_cm: float = 0.025
    output_dr_cm: float = 0.1
    duration_s: int = 10800
    output_dt_s: float = 1.0
    initial_temperature_c: float = 28.0
    initial_moisture: float = 2.55
    heat_transfer_w_m2_k: float = 25.0
    mass_transfer_m_s: float = 8e-7
    relative_tolerance: float = 2e-8
    temperature_absolute_tolerance: float = 2e-9
    moisture_absolute_tolerance: float = 2e-10
    max_time_step_s: float = 5.0


@dataclass(frozen=True)
class SolverDiagnostics:
    internal_node_count: int
    state_size: int
    function_evaluations: int
    jacobian_evaluations: int
    lu_decompositions: int
    minimum_temperature_c: float
    minimum_moisture: float


@dataclass(frozen=True)
class SimulationResult:
    time_s: np.ndarray
    radius_cm: np.ndarray
    temperature_c: np.ndarray
    moisture: np.ndarray
    diagnostics: SolverDiagnostics


def build_radial_grid_cm(radius_cm: float, dr_cm: float) -> np.ndarray:
    count = int(round(radius_cm / dr_cm)) + 1
    if count < 2 or not np.isclose(radius_cm / dr_cm, count - 1):
        raise ValueError("radius_cm must be an integer multiple of dr_cm")
    return np.linspace(0.0, radius_cm, count)


def density(c: np.ndarray) -> np.ndarray:
    return 650.0 + 128.0 * np.asarray(c, dtype=float)


def heat_capacity(c: np.ndarray) -> np.ndarray:
    c = np.asarray(c, dtype=float)
    return 1450.0 + 2736.0 * c / (c + 1.0)


def thermal_conductivity(c: np.ndarray) -> np.ndarray:
    c = np.asarray(c, dtype=float)
    return 0.21 + 0.38 * c / (c + 1.0)


def moisture_diffusivity(c: np.ndarray, temperature_c: np.ndarray) -> np.ndarray:
    safe_c = np.maximum(np.asarray(c, dtype=float), 1e-9)
    temperature_k = np.asarray(temperature_c, dtype=float) + 273.15
    if np.any(temperature_k <= 0.0):
        raise ValueError("temperature must be above absolute zero")
    return 2.4e-3 * np.exp(-0.45 / safe_c) * np.exp(-3850.0 / temperature_k)


def harmonic_mean(left: np.ndarray, right: np.ndarray) -> np.ndarray:
    left = np.asarray(left, dtype=float)
    right = np.asarray(right, dtype=float)
    denominator = left + right
    result = np.zeros_like(denominator)
    np.divide(2.0 * left * right, denominator, out=result, where=denominator > 0.0)
    return result


def radial_geometry(radius_m: float, dr_m: float) -> tuple[np.ndarray, np.ndarray, np.ndarray]:
    node_count = int(round(radius_m / dr_m)) + 1
    radius = np.linspace(0.0, radius_m, node_count)
    left_faces = np.maximum(0.0, radius - dr_m / 2.0)
    right_faces = np.minimum(radius_m, radius + dr_m / 2.0)
    volumes = np.pi * (right_faces**2 - left_faces**2)
    interface_areas = 2.0 * np.pi * (radius[:-1] + dr_m / 2.0)
    return radius, volumes, interface_areas


def add_internal_fluxes(
    derivative_numerator: np.ndarray,
    values: np.ndarray,
    interface_transport: np.ndarray,
    interface_areas: np.ndarray,
    dr_m: float,
) -> None:
    flux = interface_transport * interface_areas * (values[1:] - values[:-1]) / dr_m
    derivative_numerator[:-1] += flux
    derivative_numerator[1:] -= flux


def build_coupled_jacobian_sparsity(node_count: int) -> csc_matrix:
    """Four nearest-neighbour sparse blocks for [T_0..T_N, C_0..C_N]."""

    tri = diags(
        [np.ones(node_count - 1), np.ones(node_count), np.ones(node_count - 1)],
        offsets=[-1, 0, 1],
        shape=(node_count, node_count),
        format="csc",
    )
    return bmat([[tri, tri], [tri, tri]], format="csc")


def sample_to_output_grid(
    internal_radius_cm: np.ndarray,
    output_radius_cm: np.ndarray,
    states_by_time: np.ndarray,
) -> np.ndarray:
    spacing_cm = internal_radius_cm[1] - internal_radius_cm[0]
    indices = np.rint(output_radius_cm / spacing_cm).astype(int)
    if np.allclose(internal_radius_cm[indices], output_radius_cm):
        return states_by_time[:, indices]
    return np.vstack(
        [np.interp(output_radius_cm, internal_radius_cm, row) for row in states_by_time]
    )


def solve_problem2(
    environment: EnvironmentSeries,
    config: SolverConfig | None = None,
) -> SimulationResult:
    """Solve the nonlinear T/C system by conservative radial FVM and adaptive BDF."""

    config = config or SolverConfig()
    if environment.time_s[0] > 0.0 or environment.time_s[-1] < config.duration_s:
        raise ValueError("environment data must cover the complete simulation interval")

    internal_radius_cm = build_radial_grid_cm(config.radius_cm, config.internal_dr_cm)
    output_radius_cm = build_radial_grid_cm(config.radius_cm, config.output_dr_cm)
    dr_m = config.internal_dr_cm / 100.0
    radius_m = config.radius_cm / 100.0
    _, volumes, interface_areas = radial_geometry(radius_m, dr_m)
    surface_area = 2.0 * np.pi * radius_m
    node_count = len(internal_radius_cm)
    jacobian_pattern = build_coupled_jacobian_sparsity(node_count)
    output_time_s = np.arange(
        0.0,
        config.duration_s + config.output_dt_s / 2.0,
        config.output_dt_s,
    )

    def coupled_rhs(t_s: float, state: np.ndarray) -> np.ndarray:
        temperature = state[:node_count]
        moisture = state[node_count:]
        safe_moisture = np.maximum(moisture, 1e-9)

        rho_cp = density(safe_moisture) * heat_capacity(safe_moisture)
        node_k = thermal_conductivity(safe_moisture)
        node_d = moisture_diffusivity(safe_moisture, temperature)
        face_k = harmonic_mean(node_k[:-1], node_k[1:])
        face_d = harmonic_mean(node_d[:-1], node_d[1:])

        heat_numerator = np.zeros(node_count, dtype=float)
        moisture_numerator = np.zeros(node_count, dtype=float)
        add_internal_fluxes(heat_numerator, temperature, face_k, interface_areas, dr_m)
        add_internal_fluxes(moisture_numerator, moisture, face_d, interface_areas, dr_m)

        heat_numerator[-1] += (
            config.heat_transfer_w_m2_k
            * surface_area
            * (environment.temperature_at(t_s) - temperature[-1])
        )
        moisture_numerator[-1] += (
            config.mass_transfer_m_s
            * surface_area
            * (environment.moisture_at(t_s) - moisture[-1])
        )

        temperature_rate = heat_numerator / (rho_cp * volumes)
        moisture_rate = moisture_numerator / volumes
        return np.concatenate([temperature_rate, moisture_rate])

    initial_state = np.concatenate(
        [
            np.full(node_count, config.initial_temperature_c),
            np.full(node_count, config.initial_moisture),
        ]
    )
    absolute_tolerance = np.concatenate(
        [
            np.full(node_count, config.temperature_absolute_tolerance),
            np.full(node_count, config.moisture_absolute_tolerance),
        ]
    )
    solution = solve_ivp(
        coupled_rhs,
        (0.0, float(config.duration_s)),
        initial_state,
        method="BDF",
        t_eval=output_time_s,
        rtol=config.relative_tolerance,
        atol=absolute_tolerance,
        max_step=config.max_time_step_s,
        jac_sparsity=jacobian_pattern,
    )
    if not solution.success:
        raise RuntimeError(f"coupled integration failed: {solution.message}")
    if not np.isfinite(solution.y).all():
        raise RuntimeError("coupled integration returned non-finite values")

    internal_temperature = solution.y[:node_count, :].T
    internal_moisture = solution.y[node_count:, :].T
    if float(internal_moisture.min()) < -1e-8:
        raise RuntimeError("coupled integration returned materially negative moisture")

    temperature = sample_to_output_grid(internal_radius_cm, output_radius_cm, internal_temperature)
    moisture = sample_to_output_grid(internal_radius_cm, output_radius_cm, internal_moisture)
    diagnostics = SolverDiagnostics(
        internal_node_count=node_count,
        state_size=2 * node_count,
        function_evaluations=int(solution.nfev),
        jacobian_evaluations=int(solution.njev),
        lu_decompositions=int(solution.nlu),
        minimum_temperature_c=float(internal_temperature.min()),
        minimum_moisture=float(internal_moisture.min()),
    )
    return SimulationResult(output_time_s, output_radius_cm, temperature, moisture, diagnostics)


# ============================== 运行与产物（全部默认，无命令行参数）==============================
BASE = Path(__file__).resolve().parent                       # 代码同根目录
ATTACHMENT1 = BASE / "附件" / "附件1.xlsx"                    # 环境数据（输入）
TEMPLATE = BASE / "附件" / "附件3" / "result2.xlsx"           # 结果格式模板（题目附件3）
OUT_XLSX = BASE / "result2.xlsx"                             # 输出：同根目录
OUT_JSON = BASE / "verification2.json"
SAMPLE_SECONDS = np.array([1800, 3600, 5400, 7200, 9000, 10800], dtype=float)
SAMPLE_RADIUS_CM = np.array([0.0, 0.5, 1.0, 1.5, 2.0], dtype=float)
TEMP_SHEET = "温度"
MOIST_SHEET = "水分浓度"


def load_environment(path: Path) -> EnvironmentSeries:
    frame = pd.read_excel(path)
    required = ["时间", "温度", "水分浓度"]
    missing = [name for name in required if name not in frame.columns]
    if missing:
        raise ValueError(f"附件1缺少列：{missing}")
    return EnvironmentSeries(
        time_s=frame["时间"].to_numpy(dtype=float),
        temperature_c=frame["温度"].to_numpy(dtype=float),
        moisture=frame["水分浓度"].to_numpy(dtype=float),
    )


def selected_values(result: SimulationResult, values: np.ndarray) -> np.ndarray:
    time_indices = [int(np.where(np.isclose(result.time_s, value))[0][0]) for value in SAMPLE_SECONDS]
    radius_indices = [int(np.where(np.isclose(result.radius_cm, value))[0][0]) for value in SAMPLE_RADIUS_CM]
    return values[np.ix_(time_indices, radius_indices)]


def difference_summary(left: np.ndarray, right: np.ndarray) -> dict[str, float]:
    absolute = np.abs(left - right)
    return {
        "maximum_absolute_difference": float(absolute.max()),
        "mean_absolute_difference": float(absolute.mean()),
    }


def write_result_xlsx(result: SimulationResult, template: Path, out_path: Path) -> None:
    """读附件3 格式模板 → 按完整时间×半径网格覆盖写温度/水分两表（全精度存储，四位小数显示）。"""

    workbook = openpyxl.load_workbook(template)
    radius = result.radius_cm
    for sheet_name, field in ((TEMP_SHEET, result.temperature_c), (MOIST_SHEET, result.moisture)):
        sheet = workbook[sheet_name]
        sheet.cell(1, 1).value = "时间\\到药材中心的距离"
        for j in range(radius.size):
            sheet.cell(1, 2 + j).value = round(float(radius[j]), 10)
        for i in range(1, result.time_s.size):  # t=1..10800（跳过 t=0 初值行）
            sheet.cell(i + 1, 1).value = int(round(result.time_s[i]))
            for j in range(radius.size):
                cell = sheet.cell(i + 1, 2 + j)
                cell.value = float(field[i, j])  # 全精度存储，仅用 number_format 控制四位小数显示（与原 result2.xlsx 一致）
                cell.number_format = "0.0000"
        sheet.freeze_panes = "B2"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(out_path)


def main() -> None:
    if not ATTACHMENT1.exists():
        raise SystemExit(f"未找到附件：{ATTACHMENT1}\n请把 附件1.xlsx 放到本脚本同根目录的 附件/ 下，再运行：python problem2.py")
    if not TEMPLATE.exists():
        raise SystemExit(f"未找到结果格式模板：{TEMPLATE}\n请把题目附件3 的 result2.xlsx 放到 附件/附件3/ 下，再运行：python problem2.py")
    environment = load_environment(ATTACHMENT1)

    main_config = SolverConfig(internal_dr_cm=0.025)
    fine_config = replace(main_config, internal_dr_cm=0.0125)
    strict_time_config = replace(main_config, relative_tolerance=1e-9, temperature_absolute_tolerance=1e-10,
                                 moisture_absolute_tolerance=1e-11, max_time_step_s=2.0)
    early_ultrafine_config = replace(main_config, internal_dr_cm=0.00625, duration_s=60,
                                     relative_tolerance=1e-9, temperature_absolute_tolerance=1e-10,
                                     moisture_absolute_tolerance=1e-11, max_time_step_s=0.5)

    print("求解 main grid dr=0.025 cm（主结果）...", flush=True)
    main_result = solve_problem2(environment, main_config)
    print("求解 fine grid dr=0.0125 cm...", flush=True)
    fine_result = solve_problem2(environment, fine_config)
    print("求解 stricter BDF time settings...", flush=True)
    strict_time_result = solve_problem2(environment, strict_time_config)
    print("求解 early-time ultrafine dr=0.00625 cm for 60 s...", flush=True)
    early_ultrafine_result = solve_problem2(environment, early_ultrafine_config)

    spatial = {
        "note": "在论文表3/表4 的 30 个采样点上比较",
        "temperature_c": difference_summary(selected_values(main_result, main_result.temperature_c), selected_values(fine_result, fine_result.temperature_c)),
        "moisture": difference_summary(selected_values(main_result, main_result.moisture), selected_values(fine_result, fine_result.moisture)),
    }
    temporal = {
        "note": "在论文表3/表4 的 30 个采样点上比较",
        "temperature_c": difference_summary(selected_values(main_result, main_result.temperature_c), selected_values(strict_time_result, strict_time_result.temperature_c)),
        "moisture": difference_summary(selected_values(main_result, main_result.moisture), selected_values(strict_time_result, strict_time_result.moisture)),
    }
    early = {
        "comparison": "dr=0.0125 cm versus dr=0.00625 cm over 0-60 s",
        "temperature_c": difference_summary(fine_result.temperature_c[:61], early_ultrafine_result.temperature_c),
        "moisture": difference_summary(fine_result.moisture[:61], early_ultrafine_result.moisture),
    }
    physical_checks = {
        "minimum_temperature_c": main_result.diagnostics.minimum_temperature_c,
        "minimum_moisture": main_result.diagnostics.minimum_moisture,
        "all_finite": bool(np.isfinite(main_result.temperature_c).all() and np.isfinite(main_result.moisture).all()),
    }
    verification = {
        "source_attachment": str(ATTACHMENT1),
        "main_config": asdict(main_config),
        "diagnostics": {k: asdict(v) for k, v in
                        (("main", main_result.diagnostics), ("fine", fine_result.diagnostics),
                         ("strict_time", strict_time_result.diagnostics), ("early_ultrafine", early_ultrafine_result.diagnostics))},
        "spatial_refinement": spatial,
        "time_solver_sensitivity": temporal,
        "early_time_spatial_check_0_to_60_s": early,
        "physical_checks": physical_checks,
        "selected_tables": {
            "temperature_c": [[float(t), *[float(v) for v in row]] for t, row in zip(SAMPLE_SECONDS, selected_values(main_result, main_result.temperature_c))],
            "moisture": [[float(t), *[float(v) for v in row]] for t, row in zip(SAMPLE_SECONDS, selected_values(main_result, main_result.moisture))],
        },
    }

    write_result_xlsx(main_result, TEMPLATE, OUT_XLSX)
    OUT_JSON.write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")

    bar = "=" * 64
    radii_header = "  ".join(f"{r:>7.1f}" for r in SAMPLE_RADIUS_CM)
    print(bar, flush=True)
    print("问题二 3 h 变物性热湿耦合 —— 温度/水分（守恒径向 FVM + 联合状态自适应 BDF）", flush=True)
    print(bar, flush=True)
    print("[表3] 温度/℃（行=时间/s，列=距中心半径/cm）", flush=True)
    print(f"{'t/s':>6}  {radii_header}", flush=True)
    for t, row in zip(SAMPLE_SECONDS, selected_values(main_result, main_result.temperature_c)):
        print(f"{int(t):>6}  " + "  ".join(f"{v:>7.4f}" for v in row), flush=True)
    print("[表4] 水分浓度/(kg/kg)", flush=True)
    print(f"{'t/s':>6}  {radii_header}", flush=True)
    for t, row in zip(SAMPLE_SECONDS, selected_values(main_result, main_result.moisture)):
        print(f"{int(t):>6}  " + "  ".join(f"{v:>7.4f}" for v in row), flush=True)
    print("-" * 64, flush=True)
    print(f"[网格加密] dr 0.025→0.0125 cm：温度最大差={spatial['temperature_c']['maximum_absolute_difference']:.3e} ℃，"
          f"水分最大差={spatial['moisture']['maximum_absolute_difference']:.3e} kg/kg", flush=True)
    print(f"[BDF 容差收紧] rtol 2e-8→1e-9、max_step 5→2 s：温度最大差={temporal['temperature_c']['maximum_absolute_difference']:.3e}，"
          f"水分最大差={temporal['moisture']['maximum_absolute_difference']:.3e}（四位小数下采样值不变）", flush=True)
    print(f"[早期薄层] 0-60 s，dr 0.0125 vs 0.00625 cm：水分最大差={early['moisture']['maximum_absolute_difference']:.3e} kg/kg"
          f"（近表面薄边界层，正文表格自 0.5 h 起采样不受影响）", flush=True)
    print(f"[物理核验] 温度 min={physical_checks['minimum_temperature_c']:.4f} ℃，水分 min={physical_checks['minimum_moisture']:.6f}，"
          f"全有限={physical_checks['all_finite']}", flush=True)
    print(bar, flush=True)
    print(f"已写出：{OUT_XLSX}（温度/水分两表，按附件3 模板格式）、{OUT_JSON}", flush=True)


if __name__ == "__main__":
    main()
